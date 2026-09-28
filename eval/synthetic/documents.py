"""docx and xlsx writers.

Three past-submission layouts, named after ingestion step 3 of the plan:

- ``adjacent_cells``: one table per section; each row holds the question in one cell and the
  answer in the adjacent cell.
- ``heading_answer``: each question is a heading; its answer is the prose beneath.
- ``numbered_form``: each question is a numbered prose paragraph followed by a boxed answer
  cell (a one-column table whose header row reads "Supplier response").

Answers never contain headings, so every answer lies within one persisted section.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook

from eval.synthetic.content import (
    LAYOUT_ADJACENT_CELLS,
    LAYOUT_HEADING_ANSWER,
    LAYOUT_NUMBERED_FORM,
    REFERENCE_SECTIONS,
    REFERENCE_VARIABLES,
    ReferenceSpec,
    SubmissionSpec,
    format_date,
)


@dataclass(frozen=True)
class AnswerBlock:
    """One question-answer pair as it appears in a submission."""

    number: str
    section_title: str
    question: str
    paragraphs: tuple[str, ...]
    word_limit: int | None
    concept_id: str = ""

    @property
    def answer_text(self) -> str:
        return "\n\n".join(self.paragraphs)


@dataclass(frozen=True)
class SubmissionSection:
    number: int
    title: str
    blocks: tuple[AnswerBlock, ...]


PACK_COLUMNS: tuple[str, ...] = (
    "Section",
    "Number",
    "Question",
    "Word limit",
    "Weighting",
    "Response type",
    "Mandatory",
)

RESPONSE_TYPE_LABELS: dict[str, str] = {
    "free_text": "Free text",
    "yes_no": "Yes/No",
    "attachment": "Attachment",
    "pricing": "Pricing",
    "table": "Table",
    "other": "Other",
}


@dataclass(frozen=True)
class PackRow:
    section: str
    number: str
    question: str
    word_limit: int | None
    weighting: int | None
    response_type: str
    mandatory: bool


# ---------------------------------------------------------------------------------------------
# Past submissions
# ---------------------------------------------------------------------------------------------


def _front_matter(document: Document, spec: SubmissionSpec, context: Mapping[str, str]) -> None:
    document.add_heading(f"Response to Invitation to Tender: {spec.tender_title}", level=0)
    document.add_paragraph(f"Buyer: {spec.buyer}")
    document.add_paragraph(f"Tender reference: {spec.tender_reference}")
    document.add_paragraph(
        f"Submitted by {context['company']} on {format_date(spec.submission_date)}."
    )
    document.add_paragraph(
        f"This document is {context['company']}'s response to the invitation to tender issued "
        f"by {spec.buyer} for the {spec.tender_title}. It answers every question in the "
        f"buyer's response schedule in the order and structure the buyer set out. Pricing is "
        f"provided separately in the commercial schedule and does not appear in this document."
    )


def _fill_cell(cell, paragraphs: Sequence[str]) -> None:  # noqa: ANN001 - python-docx cell
    cell.paragraphs[0].text = paragraphs[0]
    for paragraph in paragraphs[1:]:
        cell.add_paragraph(paragraph)


def _write_adjacent_cells(document: Document, sections: Sequence[SubmissionSection]) -> None:
    for section in sections:
        document.add_heading(f"Section {section.number}: {section.title}", level=1)
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        header = table.rows[0].cells
        header[0].text = "Question"
        header[1].text = "Supplier response"
        for block in section.blocks:
            row = table.add_row().cells
            row[0].text = f"{block.number} {block.question}"
            _fill_cell(row[1], block.paragraphs)
        document.add_paragraph("")


def _write_heading_answer(document: Document, sections: Sequence[SubmissionSection]) -> None:
    for section in sections:
        document.add_heading(f"Section {section.number}: {section.title}", level=1)
        for block in section.blocks:
            document.add_heading(f"{block.number} {block.question}", level=2)
            for paragraph in block.paragraphs:
                document.add_paragraph(paragraph)


def _write_numbered_form(document: Document, sections: Sequence[SubmissionSection]) -> None:
    for section in sections:
        document.add_heading(f"Section {section.number}: {section.title}", level=1)
        for block in section.blocks:
            label = document.add_paragraph()
            label.add_run(f"Question {block.number}").bold = True
            document.add_paragraph(block.question)
            if block.word_limit:
                limit = document.add_paragraph()
                limit.add_run(f"Word limit: {block.word_limit} words").italic = True
            table = document.add_table(rows=2, cols=1)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Supplier response"
            _fill_cell(table.rows[1].cells[0], block.paragraphs)
            document.add_paragraph("")


_LAYOUT_WRITERS = {
    LAYOUT_ADJACENT_CELLS: _write_adjacent_cells,
    LAYOUT_HEADING_ANSWER: _write_heading_answer,
    LAYOUT_NUMBERED_FORM: _write_numbered_form,
}


def write_submission_docx(
    path: Path,
    spec: SubmissionSpec,
    sections: Sequence[SubmissionSection],
    context: Mapping[str, str],
) -> None:
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    _front_matter(document, spec, context)
    _LAYOUT_WRITERS[spec.layout](document, sections)
    document.save(str(path))


# ---------------------------------------------------------------------------------------------
# Reference documents
# ---------------------------------------------------------------------------------------------


def reference_first_paragraph(spec: ReferenceSpec, context: Mapping[str, str]) -> str:
    """States the effective date plainly, in the form classification is asked to extract."""
    return (
        f"This certificate confirms that the information security management system of "
        f"{context['company']} has been assessed by {context['cert_body']} and certified as "
        f"meeting the requirements of ISO/IEC 27001:2022. Certificate number "
        f"{spec.certificate_number}. Certificate issued {format_date(spec.effective_date)}. "
        f"Valid until {spec.valid_until}, subject to satisfactory surveillance audits."
    )


def write_reference_docx(path: Path, spec: ReferenceSpec, context: Mapping[str, str]) -> None:
    variables = {**context, **REFERENCE_VARIABLES[spec.key], "valid_until": spec.valid_until}
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("Certificate of Registration: ISO/IEC 27001:2022", level=0)
    document.add_paragraph(reference_first_paragraph(spec, context))
    document.add_paragraph(spec.issue_note)
    for title, sentences in REFERENCE_SECTIONS:
        document.add_heading(title, level=1)
        document.add_paragraph(" ".join(sentence.format_map(variables) for sentence in sentences))
    document.save(str(path))


# ---------------------------------------------------------------------------------------------
# Question pack
# ---------------------------------------------------------------------------------------------


def write_question_pack_xlsx(path: Path, rows: Sequence[PackRow]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questions"
    sheet.append(list(PACK_COLUMNS))
    for row in rows:
        sheet.append(
            [
                row.section,
                row.number,
                row.question,
                row.word_limit,
                row.weighting,
                RESPONSE_TYPE_LABELS[row.response_type],
                "Yes" if row.mandatory else "No",
            ]
        )
    sheet.freeze_panes = "A2"
    widths = {"A": 28, "B": 10, "C": 90, "D": 12, "E": 12, "F": 16, "G": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    workbook.save(str(path))
