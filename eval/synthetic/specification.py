"""The synthetic specification of requirements for the Northern Fells tender.

A buyer-issued Word document beside the question pack, so the Specification tab has something
to show in synthetic demo mode. It mixes the two common layouts of a real specification:
numbered requirement paragraphs under headings, and a requirements table with Ref, Requirement
and Priority columns. Its introduction carries background only and places no obligation on the
supplier, so it yields no requirements.

Each requirement records what the library should make of it (``expected``): ``covered`` when
the ingested submissions state it nearly word for word, ``partial`` when they address it in part,
and ``none`` when nothing in the library speaks to it. The wording of covered requirements
deliberately echoes the library so the demo shows a spread of suggestions; the real model judges
meaning, the synthetic double judges shared words.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.shared import Pt

SPECIFICATION_FILENAME = "specification_northern_fells_2025.docx"
SPECIFICATION_TITLE = "Specification of Requirements: Electronic Clinical Records Platform"


@dataclass(frozen=True)
class SpecRequirement:
    ref: str
    text: str
    priority: str  # must | should | could
    expected: str  # covered | partial | none


@dataclass(frozen=True)
class SpecSection:
    heading: str
    layout: str  # paragraphs | table
    requirements: tuple[SpecRequirement, ...]


INTRODUCTION: tuple[str, ...] = (
    "Northern Fells NHS Foundation Trust (fictional) provides acute and community services to a "
    "population of around 450,000 people across a largely rural area.",
    "The Trust is procuring an electronic clinical records platform to replace three legacy "
    "systems, with go-live planned in two phases over eighteen months.",
    "This document sets out the Trust's requirements. Each requirement carries a reference and a "
    "priority, and bidders state their compliance with every requirement in their response.",
)

SECTIONS: tuple[SpecSection, ...] = (
    SpecSection(
        "2. Clinical safety",
        "paragraphs",
        (
            SpecRequirement(
                "2.1",
                "The supplier must operate a clinical risk management system that meets the "
                "requirements of DCB0129 and maintain a clinical safety case for the solution.",
                "must",
                "covered",
            ),
            SpecRequirement(
                "2.2",
                "The supplier must name a Clinical Safety Officer who is a registered clinician "
                "and who approves every release before deployment.",
                "must",
                "covered",
            ),
            SpecRequirement(
                "2.3",
                "The supplier should support the Trust in discharging its own duties under "
                "DCB0160, including joint hazard workshops for local configuration and a clinical "
                "safety impact statement with every release.",
                "should",
                "partial",
            ),
        ),
    ),
    SpecSection(
        "3. Information governance and security",
        "paragraphs",
        (
            SpecRequirement(
                "3.1",
                "The supplier must hold current ISO/IEC 27001 certification covering the "
                "information security management system for the solution.",
                "must",
                "covered",
            ),
            SpecRequirement(
                "3.2",
                "The supplier must hold Cyber Essentials Plus certification covering the whole "
                "organisation and renew it annually.",
                "must",
                "covered",
            ),
            SpecRequirement(
                "3.3",
                "The supplier must have achieved Standards Met in its most recent Data Security "
                "and Protection Toolkit assessment.",
                "must",
                "covered",
            ),
            SpecRequirement(
                "3.4",
                "All personal data shall be hosted in United Kingdom public cloud regions and "
                "no personal data shall be transferred outside the United Kingdom.",
                "must",
                "covered",
            ),
            SpecRequirement(
                "3.5",
                "The supplier should commission an independent penetration test of the solution "
                "at least once a year and share a summary of the findings with the Trust.",
                "should",
                "partial",
            ),
        ),
    ),
    SpecSection(
        "4. Interoperability",
        "table",
        (
            SpecRequirement(
                "4.1",
                "The solution must expose and consume HL7 FHIR R4 resources profiled to UK Core "
                "and integrate with the Personal Demographics Service, including real-time event "
                "notifications to the Trust's integration engine.",
                "must",
                "partial",
            ),
            SpecRequirement(
                "4.2",
                "Coded clinical entries must use the SNOMED CT UK Edition, with every national "
                "terminology update applied within seven days.",
                "must",
                "partial",
            ),
            SpecRequirement(
                "4.3",
                "The solution should provide a single sign-on option using NHS Care Identity "
                "Service 2.",
                "should",
                "none",
            ),
            SpecRequirement(
                "4.4",
                "The solution could offer a patient-facing mobile application for appointment "
                "booking.",
                "could",
                "none",
            ),
        ),
    ),
    SpecSection(
        "5. Service levels and support",
        "paragraphs",
        (
            SpecRequirement(
                "5.1",
                "The supplier must provide a service desk available twenty-four hours a day, "
                "every day, by telephone, email and a customer portal.",
                "must",
                "covered",
            ),
            SpecRequirement(
                "5.2",
                "Role-based e-learning should be available to all staff before go-live, with "
                "classroom training for every clinical user.",
                "should",
                "partial",
            ),
            SpecRequirement(
                "5.3",
                "The solution must meet the Web Content Accessibility Guidelines 2.2 at level "
                "AA.",
                "must",
                "none",
            ),
        ),
    ),
    SpecSection(
        "6. Social value",
        "paragraphs",
        (
            SpecRequirement(
                "6.1",
                "The supplier must have published a Carbon Reduction Plan in the format required "
                "by Procurement Policy Note 06/21, with a commitment to net zero.",
                "must",
                "covered",
            ),
            SpecRequirement(
                "6.2",
                "The supplier should fund at least five apprenticeships for young people from the "
                "Trust's area in every contract year.",
                "should",
                "partial",
            ),
        ),
    ),
)

PRIORITY_LABELS = {"must": "Must", "should": "Should", "could": "Could"}


def all_requirements() -> list[SpecRequirement]:
    return [requirement for section in SECTIONS for requirement in section.requirements]


def write_specification_docx(path: Path) -> None:
    """Write the specification: background, numbered paragraphs and a requirements table."""
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(SPECIFICATION_TITLE, level=0)
    document.add_heading("1. Introduction", level=1)
    for paragraph in INTRODUCTION:
        document.add_paragraph(paragraph)
    for section in SECTIONS:
        document.add_heading(section.heading, level=1)
        if section.layout == "table":
            table = document.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for cell, title in zip(table.rows[0].cells, ("Ref", "Requirement", "Priority"),
                                   strict=True):
                cell.text = title
            for requirement in section.requirements:
                cells = table.add_row().cells
                cells[0].text = requirement.ref
                cells[1].text = requirement.text
                cells[2].text = PRIORITY_LABELS[requirement.priority]
        else:
            for requirement in section.requirements:
                document.add_paragraph(f"{requirement.ref} {requirement.text}")
    document.save(str(path))


def manifest_entry(buyer: str) -> dict[str, object]:
    return {
        "filename": SPECIFICATION_FILENAME,
        "role": "specification",
        "expected": {
            "doc_type": "tender_document",
            "tender_doc_kind": "specification",
            "buyer": buyer,
            "requirement_count": len(all_requirements()),
        },
        "requirements": [
            {
                "ref": requirement.ref,
                "priority": requirement.priority,
                "text": requirement.text,
                "expected": requirement.expected,
            }
            for requirement in all_requirements()
        ],
    }


__all__ = [
    "SPECIFICATION_FILENAME",
    "SpecRequirement",
    "all_requirements",
    "manifest_entry",
    "write_specification_docx",
]
