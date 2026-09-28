"""The ``extract_questions`` job: xlsx and docx packs built in the test, a section parser
standing in for owner A's Docling path, scripted FakeLLM extraction, and the handler run
through the job registry and worker loop."""

from __future__ import annotations

import re
import uuid
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import DEFAULT_ORG_ID, get_settings
from app.db.models import Document, DocumentSection, Job, Question, Tender, Thread
from app.ingest.types import ParsedSection
from app.jobs import enqueue
from app.llm.embeddings import embed
from app.worker import run_job

# Imported by the autouse fixture below rather than at module level: importing the module
# registers the extract_questions handler, and registration belongs to the test run, not to
# collection (see the note in test_triage_jobs.py).
q: Any = None


@pytest.fixture(autouse=True)
def _load_questions_module() -> None:
    global q
    from app.ingest import questions as module

    q = module

SECTIONS = ("1. Clinical Safety", "2. Information Governance", "3. Commercial")
ROW_COUNT = 30  # two row batches: 25 + 5


# --- Building packs -----------------------------------------------------------------------------


def build_xlsx_pack(path: Path, *, rows: int = ROW_COUNT) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questions"
    sheet.append(["Ref", "Section", "Question", "Word limit", "Weighting", "Mandatory"])
    for index in range(1, rows + 1):
        section = SECTIONS[(index - 1) % len(SECTIONS)]
        sheet.append(
            [
                f"Q{index}",
                section,
                f"Describe how your organisation handles topic {index} for the buyer.",
                500 if index % 2 else None,
                10 if index % 3 else None,
                "Yes" if index % 5 == 0 else "No",
            ]
        )
    workbook.save(path)
    return path


def build_docx_pack(path: Path) -> Path:
    document = DocxDocument()
    document.add_paragraph("Instructions to bidders: answer every question in full.")
    document.add_heading("3. Clinical Safety", level=1)
    document.add_paragraph(
        "3.1 Describe your clinical safety case and name your clinical safety officer. "
        "(Maximum 500 words, weighting 15%)"
    )
    document.add_paragraph("3.2 Confirm compliance with DCB0129. Mandatory: yes.")
    document.add_heading("4. Social Value", level=1)
    document.add_paragraph("4.1 Outline your social value commitments. (Maximum 300 words)")
    document.save(path)
    return path


# --- A stand-in parser at the plan's section granularity -----------------------------------------


def parse_stand_in(path: Path, filename: str) -> list[ParsedSection]:
    if filename.endswith(".xlsx"):
        workbook = load_workbook(path, read_only=True)
        sections: list[ParsedSection] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(cell) if cell is not None else "" for cell in rows[0]]
            for row_number, row in enumerate(rows[1:], start=2):
                cells = ["" if cell is None else str(cell) for cell in row]
                if not any(cells):
                    continue
                sections.append(
                    ParsedSection(
                        order_index=len(sections),
                        heading_path=[sheet.title, *header],
                        page_start=None,
                        page_end=None,
                        table_index=0,
                        cell_ref=f"{sheet.title}!{row_number}",
                        text=get_settings().table_cell_delimiter.join(cells),
                    )
                )
        return sections
    document = DocxDocument(str(path))
    sections = []
    heading_path: list[str] = []
    current: list[str] = []

    def flush() -> None:
        if current:
            sections.append(
                ParsedSection(
                    order_index=len(sections),
                    heading_path=list(heading_path),
                    page_start=1,
                    page_end=1,
                    table_index=None,
                    cell_ref=None,
                    text="\n".join(current),
                )
            )
            current.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.startswith("Heading"):
            flush()
            heading_path = [text]
            current.append(text)
        else:
            current.append(text)
    flush()
    return sections


def persist_stand_in(
    session: Session, document: Document, parsed: Sequence[ParsedSection]
) -> list[DocumentSection]:
    rows = [
        DocumentSection(
            org_id=document.org_id,
            document_id=document.id,
            order_index=section.order_index,
            heading_path=section.heading_path,
            page_start=section.page_start,
            page_end=section.page_end,
            table_index=section.table_index,
            cell_ref=section.cell_ref,
            text=section.text,
        )
        for section in parsed
    ]
    session.add_all(rows)
    session.flush()
    return rows


@pytest.fixture
def stand_in_parser(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(q, "_load_parser", lambda: (parse_stand_in, persist_stand_in))


# --- Scripted extraction ------------------------------------------------------------------------

_SECTION_BLOCK = re.compile(r"\[section ([0-9a-f-]{36})\] \((prose|table row)[^)]*\)\n")
_NUMBERED = re.compile(r"^(\d+\.\d+)\s+(.+)$", re.MULTILINE)


def scripted_extraction(*, user: str, **_: object) -> dict:
    """Read the sections out of the prompt and answer as a careful model would."""
    taxonomy_line = next(line for line in user.splitlines() if line.startswith("Topic taxonomy"))
    taxonomy = [t.strip() for t in taxonomy_line.split(":", 1)[1].split(",")]
    blocks = _SECTION_BLOCK.split(user)
    questions: list[dict] = []
    # blocks = [preamble, id, kind, body, id, kind, body, ...]
    for index in range(1, len(blocks), 3):
        kind, body = blocks[index + 1], blocks[index + 2]
        text = body.split("\n", 1)[1] if "\n" in body else body
        if kind == "table row":
            cells = [cell.strip() for cell in text.strip().split(" | ")]
            ref, section, question, limit, weighting, mandatory = cells[:6]
            questions.append(
                {
                    "section": section,
                    "number": ref,
                    "text": question,
                    "word_limit": int(limit) if limit else None,
                    "weighting": float(weighting) if weighting else None,
                    "response_type": "pricing" if "Commercial" in section else "free_text",
                    "mandatory": mandatory.lower() == "yes",
                    "order_index": len(questions),
                    "topics": [
                        "commercial_and_pricing" if "Commercial" in section else taxonomy[0],
                        "not_a_topic",
                    ],
                }
            )
        else:
            heading = text.strip().splitlines()[0]
            for match in _NUMBERED.finditer(text):
                body_text = match.group(2)
                limit = re.search(r"Maximum (\d+) words", body_text)
                weight = re.search(r"weighting (\d+)%", body_text)
                questions.append(
                    {
                        "section": heading,
                        "number": match.group(1),
                        "text": body_text,
                        "word_limit": int(limit.group(1)) if limit else None,
                        "weighting": float(weight.group(1)) if weight else None,
                        "response_type": "yes_no" if "Confirm" in body_text else "free_text",
                        "mandatory": "Mandatory: yes" in body_text,
                        "order_index": len(questions),
                        "topics": ["social_value" if "social value" in body_text.lower()
                                   else "clinical_safety"],
                    }
                )
    return {"questions": questions}


# --- Helpers ------------------------------------------------------------------------------------


def make_tender(session: Session, name: str = "Community diagnostics") -> Tender:
    tender = Tender(org_id=DEFAULT_ORG_ID, name=name, buyer="NHS Test ICB")
    session.add(tender)
    session.flush()
    return tender


def make_pack(session: Session, tender: Tender, path: Path) -> Document:
    document = Document(
        org_id=DEFAULT_ORG_ID,
        filename=path.name,
        storage_path=str(path),
        doc_type="tender_document",
        tender_id=tender.id,
        tender_doc_kind="question_pack",
        classification_confirmed=True,
        ingest_status="queued",
    )
    session.add(document)
    session.flush()
    return document


def enqueue_extract(session: Session, tender: Tender, document: Document) -> Job:
    job = enqueue(
        session,
        "extract_questions",
        {"tender_id": str(tender.id), "document_id": str(document.id), "actor": "test user"},
    )
    tender.extract_job_id = job.id
    session.flush()
    return job


def run(session: Session, job: Job) -> bool:
    """Claim this job the way the worker would and dispatch it through the registry."""
    job.status = "running"
    job.attempts += 1
    session.flush()
    run_job(session, job)
    return True


# --- Tests --------------------------------------------------------------------------------------


def test_build_batches_prose_singly_and_rows_in_batches_of_25() -> None:
    def section(index: int, *, row: bool, text: str = "x") -> DocumentSection:
        return DocumentSection(
            order_index=index,
            heading_path=[],
            table_index=0 if row else None,
            cell_ref=f"S!{index}" if row else None,
            text=text,
        )

    sections = [section(0, row=False), *(section(i, row=True) for i in range(1, 31))]
    sections.append(section(31, row=False, text="   "))  # empty prose is skipped
    sections.append(section(32, row=False))
    batches = q.build_batches(sections)
    assert [len(batch.sections) for batch in batches] == [1, 25, 5, 1]
    assert [batch.is_rows for batch in batches] == [False, True, True, False]


def test_xlsx_pack_end_to_end_through_the_registry(
    db_session: Session, tmp_path: Path, fake_llm, fake_embeddings, stand_in_parser  # noqa: ANN001
) -> None:
    fake_llm.register("extract_questions", scripted_extraction)
    tender = make_tender(db_session)
    pack = make_pack(db_session, tender, build_xlsx_pack(tmp_path / "pack.xlsx"))
    job = enqueue_extract(db_session, tender, pack)

    assert run(db_session, job) is True

    db_session.refresh(job)
    db_session.refresh(pack)
    db_session.refresh(tender)
    assert job.status == "done", job.error
    assert pack.ingest_status == "ready"
    assert job.total == 2 and job.done == 2, "25 rows in the first batch, 5 in the second"
    assert [r["outcome"] for r in job.results] == ["extracted", "extracted"]
    assert [r["detail"] for r in job.results] == [25, 5]
    first_section = db_session.scalars(
        select(DocumentSection)
        .where(DocumentSection.document_id == pack.id)
        .order_by(DocumentSection.order_index)
    ).first()
    assert job.results[0]["item_id"] == str(first_section.id)
    assert len(fake_llm.calls) == 2
    assert fake_llm.calls[0].model == get_settings().model_main
    assert "clinical_safety" in fake_llm.calls[0].user, "taxonomy is passed to the model"

    questions = db_session.scalars(
        select(Question).where(Question.tender_id == tender.id).order_by(Question.order_index)
    ).all()
    assert len(questions) == ROW_COUNT
    assert [question.order_index for question in questions] == list(range(ROW_COUNT))
    assert {question.status for question in questions} == {"not_started"}
    assert {question.coverage for question in questions} == {"unknown"}
    assert all(question.document_id == pack.id for question in questions)
    first = questions[0]
    assert (first.section, first.number) == ("1. Clinical Safety", "Q1")
    assert first.word_limit == 500 and first.weighting == 10.0 and first.mandatory is False
    assert first.topics == ["clinical_safety"], "unknown taxonomy ids are dropped"
    commercial = next(question for question in questions if question.number == "Q3")
    assert commercial.response_type == "pricing"
    assert commercial.topics == ["commercial_and_pricing"]
    assert next(question for question in questions if question.number == "Q5").mandatory is True
    assert next(question for question in questions if question.number == "Q2").word_limit is None

    # One thread per question, titled by section and number.
    threads = db_session.scalars(select(Thread).where(Thread.tender_id == tender.id)).all()
    assert sorted(thread.question_id for thread in threads) == sorted(x.id for x in questions)
    assert any(thread.title == "1. Clinical Safety Q1" for thread in threads)

    # Every question text embedded, matching a direct call to embed().
    assert all(question.embedding is not None for question in questions)
    assert list(first.embedding) == pytest.approx(embed([first.text])[0])

    # Chaining: triage enqueued, recorded on this job and on the tender.
    assert job.next_job_id is not None
    assert tender.triage_job_id == job.next_job_id
    triage = db_session.get(Job, job.next_job_id)
    assert triage.kind == "triage_tender" and triage.status == "queued"
    assert triage.payload == {"tender_id": str(tender.id), "actor": "test user", "scope": "unknown"}
    assert triage.total == ROW_COUNT


def test_rerun_is_idempotent_on_section_and_number(
    db_session: Session, tmp_path: Path, fake_llm, fake_embeddings, stand_in_parser  # noqa: ANN001
) -> None:
    fake_llm.register("extract_questions", scripted_extraction)
    tender = make_tender(db_session)
    pack = make_pack(db_session, tender, build_xlsx_pack(tmp_path / "pack.xlsx", rows=6))
    first = enqueue_extract(db_session, tender, pack)
    assert run(db_session, first)
    before = {
        (question.section, question.number): question.id
        for question in db_session.scalars(select(Question).where(Question.tender_id == tender.id))
    }
    assert len(before) == 6

    # Rewrite the pack: same references, one changed text; run again.
    workbook = load_workbook(pack.storage_path)
    workbook["Questions"]["C2"] = "Describe your revised approach to clinical safety."
    workbook.save(pack.storage_path)
    # The sections were persisted once and are reused, so the change must reach the model
    # through a fresh parse in a real re-upload; here we edit the stored row to mirror that.
    section = db_session.scalars(
        select(DocumentSection).where(DocumentSection.document_id == pack.id).order_by(
            DocumentSection.order_index
        )
    ).first()
    section.text = section.text.replace(
        "Describe how your organisation handles topic 1 for the buyer.",
        "Describe your revised approach to clinical safety.",
    )
    db_session.flush()
    fake_llm.calls.clear()
    second = enqueue_extract(db_session, tender, pack)
    assert run(db_session, second)
    db_session.refresh(second)
    assert second.status == "done", second.error

    after = {
        (question.section, question.number): question
        for question in db_session.scalars(select(Question).where(Question.tender_id == tender.id))
    }
    assert len(after) == 6, "no duplicates on re-run"
    assert {key: question.id for key, question in after.items()} == before
    assert after[("1. Clinical Safety", "Q1")].text == (
        "Describe your revised approach to clinical safety."
    )
    assert len(fake_llm.calls) == 1, "sections are reused, not re-parsed"
    threads = db_session.scalars(select(Thread).where(Thread.tender_id == tender.id)).all()
    assert len(threads) == 6, "threads are not duplicated either"


def test_docx_pack_prose_sections_one_call_each(
    db_session: Session, tmp_path: Path, fake_llm, fake_embeddings, stand_in_parser  # noqa: ANN001
) -> None:
    fake_llm.register("extract_questions", scripted_extraction)
    tender = make_tender(db_session)
    pack = make_pack(db_session, tender, build_docx_pack(tmp_path / "pack.docx"))
    job = enqueue_extract(db_session, tender, pack)
    assert run(db_session, job)
    db_session.refresh(job)
    assert job.status == "done", job.error
    # Preamble, "3. Clinical Safety" block, "4. Social Value" block: three prose sections.
    assert job.total == 3 and len(fake_llm.calls) == 3
    assert [r["detail"] for r in job.results] == [0, 2, 1]

    questions = db_session.scalars(
        select(Question).where(Question.tender_id == tender.id).order_by(Question.order_index)
    ).all()
    assert [(x.section, x.number) for x in questions] == [
        ("3. Clinical Safety", "3.1"),
        ("3. Clinical Safety", "3.2"),
        ("4. Social Value", "4.1"),
    ]
    assert questions[0].word_limit == 500 and questions[0].weighting == 15.0
    assert questions[1].response_type == "yes_no" and questions[1].mandatory is True
    assert questions[2].topics == ["social_value"]


def test_batch_failure_is_recorded_and_does_not_fail_the_job(
    db_session: Session, tmp_path: Path, fake_llm, fake_embeddings, stand_in_parser  # noqa: ANN001
) -> None:
    calls = {"n": 0}

    def flaky(**context: object) -> dict:
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("model unavailable")
        return scripted_extraction(**context)  # type: ignore[arg-type]

    fake_llm.register("extract_questions", flaky)
    tender = make_tender(db_session)
    pack = make_pack(db_session, tender, build_xlsx_pack(tmp_path / "pack.xlsx"))
    job = enqueue_extract(db_session, tender, pack)
    assert run(db_session, job)
    db_session.refresh(job)
    assert job.status == "done"
    assert [r["outcome"] for r in job.results] == ["failed", "extracted"]
    assert "model unavailable" in job.results[0]["detail"]
    assert job.results[1]["detail"] == 5
    questions = db_session.scalars(select(Question).where(Question.tender_id == tender.id)).all()
    assert len(questions) == 5
    assert [x.order_index for x in sorted(questions, key=lambda x: x.order_index)] == list(range(5))
    assert job.next_job_id is not None, "triage is still chained"


def test_missing_parser_fails_the_attempt_with_a_clear_error(
    db_session: Session, tmp_path: Path, fake_llm, monkeypatch: pytest.MonkeyPatch  # noqa: ANN001
) -> None:
    def unavailable() -> tuple:
        raise q.ExtractionUnavailable(
            "Question-pack parsing is unavailable: app.ingest.parse could not be imported."
        )

    monkeypatch.setattr(q, "_load_parser", unavailable)
    tender = make_tender(db_session)
    pack = make_pack(db_session, tender, build_xlsx_pack(tmp_path / "pack.xlsx"))
    job = enqueue_extract(db_session, tender, pack)
    assert run(db_session, job)
    db_session.refresh(job)
    db_session.refresh(pack)
    assert job.status == "queued", "below max_attempts the worker requeues"
    assert "app.ingest.parse" in (job.error or "")
    assert pack.ingest_status == "failed"
    assert fake_llm.calls == []


def test_normalise_questions_fills_section_and_number_and_caps_topics() -> None:
    section = DocumentSection(
        id=uuid.uuid4(), order_index=0, heading_path=["Lot 2", "Ref", "Question"], text="row"
    )
    batch = q.ExtractionBatch([section])
    output = q.QuestionExtractionOutput(
        questions=[
            q.ExtractedQuestion(text="Second", order_index=1, topics=["a", "b"]),
            q.ExtractedQuestion(
                text="First",
                order_index=0,
                word_limit=0,
                topics=["clinical_safety", "Social_Value", "interoperability", "service_levels"],
            ),
            q.ExtractedQuestion(text="   ", order_index=2),
        ]
    )
    items = q.normalise_questions(
        output, batch, get_settings().default_topic_taxonomy, position_offset=10
    )
    assert [item.text for item in items] == ["First", "Second"]
    assert items[0].section == "Lot 2" and items[0].number == "11"
    assert items[1].number == "12"
    assert items[0].word_limit is None, "a zero limit is no limit"
    assert items[0].topics == ["clinical_safety", "social_value", "interoperability"]
    assert items[1].topics == []


@pytest.mark.integration
def test_real_parser_on_xlsx_pack(
    db_session: Session, tmp_path: Path, fake_llm, fake_embeddings  # noqa: ANN001
) -> None:
    """Owner A's parser, when present, must yield one section per spreadsheet row."""
    pytest.importorskip("app.ingest.parse")
    fake_llm.register("extract_questions", scripted_extraction)
    tender = make_tender(db_session)
    pack = make_pack(db_session, tender, build_xlsx_pack(tmp_path / "pack.xlsx", rows=4))
    job = enqueue_extract(db_session, tender, pack)
    assert run(db_session, job)
    db_session.refresh(job)
    assert job.status == "done", job.error
    sections = db_session.scalars(
        select(DocumentSection).where(DocumentSection.document_id == pack.id)
    ).all()
    assert len(sections) == 4
    assert all(section.cell_ref for section in sections)
