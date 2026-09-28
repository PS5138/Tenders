"""xlsx export: header row, one row per question in order, placeholder and status labels."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.export import PLACEHOLDER, ExportBlocked, build_export
from app.export.xlsx import HEADERS, SHEET_TITLE
from tests.test_export_helpers import make_answer, make_question, make_question_pack, make_tender


def _rows(payload: bytes) -> list[tuple[object, ...]]:
    workbook = load_workbook(io.BytesIO(payload))
    sheet = workbook[SHEET_TITLE]
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def test_submission_xlsx_one_row_per_question(db_session: Session) -> None:
    tender = make_tender(db_session, name="Community EPR", buyer="NHS Test ICB")
    pack = make_question_pack(db_session, tender)
    approved = make_question(
        db_session, tender, pack, order_index=2, section="B", number="2.1",
        text="Approved question?", status="approved",
    )
    make_question(
        db_session, tender, pack, order_index=1, section="A", number="1.1",
        text="Drafted question?", status="ai_draft",
    )
    make_answer(db_session, approved, text="Four words in here.")

    rows = _rows(build_export(db_session, tender, format="xlsx", mode="submission"))

    assert rows[0] == HEADERS
    assert rows[1] == ("A", "1.1", "Drafted question?", PLACEHOLDER, "AI draft", None)
    assert rows[2] == ("B", "2.1", "Approved question?", "Four words in here.", "Approved", 4)
    assert len(rows) == 3

    payload = build_export(db_session, tender, format="xlsx", mode="submission")
    workbook = load_workbook(io.BytesIO(payload))
    assert workbook.properties.title == "Community EPR"
    assert workbook.properties.subject == "NHS Test ICB"
    assert workbook[SHEET_TITLE].freeze_panes == "A2"


def test_review_xlsx_writes_current_answers_regardless_of_status(db_session: Session) -> None:
    tender = make_tender(db_session)
    pack = make_question_pack(db_session, tender)
    drafted = make_question(
        db_session, tender, pack, order_index=1, section="A", number="1", text="Q1?",
        status="writer_edited", needs_review=True,
    )
    make_question(db_session, tender, pack, order_index=2, section="A", number="2", text="Q2?")
    make_answer(db_session, drafted, text="Edited by a person.", author_type="user")

    rows = _rows(build_export(db_session, tender, format="xlsx", mode="review"))
    assert rows[1] == ("A", "1", "Q1?", "Edited by a person.", "Writer edited", 4)
    assert rows[2] == ("A", "2", "Q2?", PLACEHOLDER, "Not started", None)

    with pytest.raises(ExportBlocked):
        build_export(db_session, tender, format="xlsx", mode="submission")
