"""Tests for the synthetic data generator (``eval/generate_synthetic.py``), offline mode.

Generation runs once per module into a temporary directory with the fake LLM provider; the
files are then opened with python-docx and openpyxl and checked against the plan's
requirements for the demo data.
"""

from __future__ import annotations

import difflib
import json
import sys
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from eval.generate_synthetic import generate, main  # noqa: E402
from eval.synthetic import content  # noqa: E402
from eval.synthetic.documents import PACK_COLUMNS  # noqa: E402

# ---------------------------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------------------------


def docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def cell_text(cell) -> str:  # noqa: ANN001 - python-docx cell
    """Cell paragraphs joined the way the other layouts join theirs."""
    return "\n\n".join(p.text for p in cell.paragraphs if p.text.strip())


def answers_by_concept(out: Path, doc: dict) -> dict[str, str]:
    """Recover each pair's answer text from the docx, following its layout."""
    document = Document(str(out / doc["filename"]))
    pairs = doc["pairs"]
    if doc["layout"] == content.LAYOUT_ADJACENT_CELLS:
        rows = [row for table in document.tables for row in table.rows[1:]]
        assert len(rows) == len(pairs)
        return {
            pair["concept_id"]: cell_text(row.cells[1])
            for pair, row in zip(pairs, rows, strict=True)
        }
    if doc["layout"] == content.LAYOUT_NUMBERED_FORM:
        assert len(document.tables) == len(pairs)
        return {
            pair["concept_id"]: cell_text(table.rows[1].cells[0])
            for pair, table in zip(pairs, document.tables, strict=True)
        }
    texts: dict[str, list[str]] = {}
    current: str | None = None
    index = -1
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Heading 2":
            index += 1
            current = pairs[index]["concept_id"]
            texts[current] = []
        elif current and paragraph.style.name == "Normal" and paragraph.text.strip():
            texts[current].append(paragraph.text)
    assert len(texts) == len(pairs)
    return {concept_id: "\n\n".join(paragraphs) for concept_id, paragraphs in texts.items()}


def ratio(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()


# ---------------------------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------------------------


@pytest.fixture(scope="module")
def out(tmp_path_factory: pytest.TempPathFactory) -> Path:
    directory = tmp_path_factory.mktemp("synthetic")
    generate(directory, seed=0, mode="template")
    return directory


@pytest.fixture(scope="module")
def manifest(out: Path) -> dict:
    return json.loads((out / "manifest.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def ground_truth(out: Path) -> dict:
    return json.loads((out / "ground_truth.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def submissions(manifest: dict) -> dict[str, dict]:
    return {
        doc["key"]: doc
        for doc in manifest["documents"]
        if doc["role"] in ("past_submission", "held_out_submission")
    }


@pytest.fixture(scope="module")
def pack_rows(out: Path) -> list[tuple]:
    sheet = load_workbook(out / content.QUESTION_PACK_FILENAME, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert tuple(rows[0]) == PACK_COLUMNS
    return rows[1:]


# ---------------------------------------------------------------------------------------------
# Files and modes
# ---------------------------------------------------------------------------------------------


def test_every_output_file_exists(out: Path, manifest: dict) -> None:
    filenames = {doc["filename"] for doc in manifest["documents"]}
    assert len([f for f in filenames if f.endswith(".docx")]) == 5
    assert len([f for f in filenames if f.endswith(".xlsx")]) == 1
    for filename in filenames:
        assert (out / filename).is_file(), filename
    assert (out / "ground_truth.json").is_file()
    assert manifest["mode"] == "template"


def test_offline_mode_makes_no_llm_calls(tmp_path: Path, fake_llm) -> None:  # noqa: ANN001
    manifest = generate(tmp_path, seed=3)  # mode resolved from LLM_PROVIDER=fake
    assert manifest["mode"] == "template"
    assert fake_llm.calls == []


def test_cli_entry_point(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["--out", str(tmp_path), "--seed", "1", "--mode", "template"]) == 0
    assert "Wrote 6 documents" in capsys.readouterr().out
    assert (tmp_path / "ground_truth.json").is_file()


def test_generation_is_deterministic(tmp_path: Path) -> None:
    first, second = tmp_path / "a", tmp_path / "b"
    generate(first, seed=0, mode="template")
    generate(second, seed=0, mode="template")
    assert (first / "ground_truth.json").read_text() == (second / "ground_truth.json").read_text()
    for filename in [spec.filename for spec in content.SUBMISSIONS]:
        assert docx_text(first / filename) == docx_text(second / filename)


def test_seed_changes_prose_but_not_facts(tmp_path: Path) -> None:
    generate(tmp_path / "s0", seed=0, mode="template")
    generate(tmp_path / "s7", seed=7, mode="template")
    spec = content.SUBMISSIONS[0]
    a, b = docx_text(tmp_path / "s0" / spec.filename), docx_text(tmp_path / "s7" / spec.filename)
    assert a != b
    snapshot = content.facts_as_at(spec.submission_date)
    for text in (a, b):
        assert snapshot.iso_cert in text and snapshot.iso_issued in text


# ---------------------------------------------------------------------------------------------
# Past submissions: layouts, counts, no headings inside answers
# ---------------------------------------------------------------------------------------------


def test_the_three_layouts_are_all_used(submissions: dict[str, dict]) -> None:
    assert {doc["layout"] for doc in submissions.values()} == set(content.LAYOUTS)
    assert sum(1 for doc in submissions.values() if doc["role"] == "held_out_submission") == 1


def test_adjacent_cells_layout(out: Path, submissions: dict[str, dict]) -> None:
    doc = next(d for d in submissions.values() if d["layout"] == content.LAYOUT_ADJACENT_CELLS)
    document = Document(str(out / doc["filename"]))
    assert document.tables, "adjacent-cells layout needs tables"
    rows = 0
    for table in document.tables:
        assert len(table.columns) == 2
        assert [cell.text for cell in table.rows[0].cells] == ["Question", "Supplier response"]
        for row in table.rows[1:]:
            question, answer = (cell.text for cell in row.cells)
            assert question.strip() and answer.strip()
            rows += 1
    assert rows == len(doc["pairs"])
    # Question and answer are in the same row: the question text sits beside its answer.
    first = doc["pairs"][0]
    assert first["question_text"] in document.tables[0].rows[1].cells[0].text


def test_heading_answer_layout(out: Path, submissions: dict[str, dict]) -> None:
    doc = next(d for d in submissions.values() if d["layout"] == content.LAYOUT_HEADING_ANSWER)
    document = Document(str(out / doc["filename"]))
    assert not document.tables
    headings = [p for p in document.paragraphs if p.style.name == "Heading 2"]
    assert len(headings) == len(doc["pairs"])
    for heading, pair in zip(headings, doc["pairs"], strict=True):
        assert pair["question_text"] in heading.text
    # Every heading is followed by at least one Normal paragraph before the next heading.
    styles = [p.style.name for p in document.paragraphs if p.text.strip()]
    for index, style in enumerate(styles):
        if style == "Heading 2":
            assert styles[index + 1] == "Normal"


def test_numbered_form_layout(out: Path, submissions: dict[str, dict]) -> None:
    doc = next(d for d in submissions.values() if d["layout"] == content.LAYOUT_NUMBERED_FORM)
    document = Document(str(out / doc["filename"]))
    assert len(document.tables) == len(doc["pairs"])
    for table in document.tables:
        assert len(table.columns) == 1 and len(table.rows) == 2
        assert table.rows[0].cells[0].text == "Supplier response"
        assert table.rows[1].cells[0].text.strip()
    body = "\n".join(p.text for p in document.paragraphs)
    for pair in doc["pairs"]:
        assert f"Question {pair['number']}" in body
        assert pair["question_text"] in body


def test_each_submission_has_ten_to_fourteen_pairs(
    out: Path, submissions: dict[str, dict]
) -> None:
    for doc in submissions.values():
        assert 10 <= len(doc["pairs"]) <= 14, doc["filename"]
        assert doc["expected"]["pair_count"] == len(doc["pairs"])
        assert len(answers_by_concept(out, doc)) == len(doc["pairs"])


def test_answers_contain_no_headings(out: Path, submissions: dict[str, dict]) -> None:
    for doc in submissions.values():
        document = Document(str(out / doc["filename"]))
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    assert all(p.style.name == "Normal" for p in cell.paragraphs), doc["filename"]
        for answer in answers_by_concept(out, doc).values():
            for line in answer.splitlines():
                assert not line.startswith(("#", "*", "-")), line
                assert line.strip() == "" or line.rstrip().endswith((".", "!", "?")), line


def test_answers_are_prose_of_reasonable_length(out: Path, submissions: dict[str, dict]) -> None:
    for doc in submissions.values():
        for concept_id, answer in answers_by_concept(out, doc).items():
            words = len(answer.split())
            limit = content.CONCEPTS[concept_id].word_limit or 500
            assert 50 <= words <= limit, (doc["filename"], concept_id, words)
            assert answer.count("\n\n") >= 1, "answers have at least two paragraphs"


def test_submission_front_matter_states_buyer_and_date(
    out: Path, submissions: dict[str, dict]
) -> None:
    for spec in content.SUBMISSIONS:
        text = docx_text(out / spec.filename)
        assert spec.buyer in text
        assert content.format_date(spec.submission_date) in text
        assert content.COMPANY["company"] in text


def test_dated_facts_appear_as_single_sentences(out: Path, submissions: dict[str, dict]) -> None:
    """The fact extractor copies one sentence per fact; the certificate sentences must be there."""
    for spec in content.SUBMISSIONS:
        snapshot = content.facts_as_at(spec.submission_date)
        text = docx_text(out / spec.filename)
        if "is_iso27001" in spec.concept_ids:
            assert (
                f"number {snapshot.iso_cert}, was issued on {snapshot.iso_issued} by "
                f"{content.COMPANY['cert_body']}"
            ) in text
        if "is_ce_plus" in spec.concept_ids:
            assert f"number {snapshot.ce_cert}, was issued on {snapshot.ce_issued}" in text
            assert f"expires on {snapshot.ce_expires}" in text
        if "ig_dspt" in spec.concept_ids:
            assert f"{snapshot.dspt_year} Data Security and Protection Toolkit" in text
            assert snapshot.dspt_published in text
        if "cs_officer" in spec.concept_ids:
            assert f"Our Clinical Safety Officer is {content.COMPANY['cso']}" in text
        if "ig_dpo" in spec.concept_ids:
            assert f"Our Data Protection Officer is {content.COMPANY['dpo']}" in text


# ---------------------------------------------------------------------------------------------
# Overlap: deduplication and the held-out counterparts
# ---------------------------------------------------------------------------------------------


def test_near_identical_answers_across_ingested_submissions(
    out: Path, submissions: dict[str, dict]
) -> None:
    ingested = [d for d in submissions.values() if d["role"] == "past_submission"]
    assert len(ingested) == 2
    texts = {doc["key"]: answers_by_concept(out, doc) for doc in ingested}
    overlaps = [
        (doc["key"], link["concept_id"], link["submission_filename"])
        for doc in ingested
        for link in doc["near_identical_to"]
    ]
    assert len(overlaps) >= 3, "several questions reappear across submissions"
    by_filename = {doc["filename"]: doc["key"] for doc in ingested}
    for key, concept_id, source_filename in overlaps:
        source_key = by_filename[source_filename]
        a, b = texts[key][concept_id], texts[source_key][concept_id]
        assert a != b
        assert ratio(a, b) >= 0.8, (concept_id, ratio(a, b))


def test_shared_concepts_reuse_the_question_in_different_words(
    submissions: dict[str, dict],
) -> None:
    ingested = [d for d in submissions.values() if d["role"] == "past_submission"]
    questions = {
        doc["key"]: {pair["concept_id"]: pair["question_text"] for pair in doc["pairs"]}
        for doc in ingested
    }
    first, second = questions.values()
    shared = set(first) & set(second)
    assert len(shared) >= 3
    assert any(first[c] != second[c] for c in shared)


def test_held_out_questions_have_counterparts(submissions: dict[str, dict]) -> None:
    held_out = next(d for d in submissions.values() if d["role"] == "held_out_submission")
    ingested_concepts = {
        pair["concept_id"]
        for doc in submissions.values()
        if doc["role"] == "past_submission"
        for pair in doc["pairs"]
    }
    held_out_concepts = [pair["concept_id"] for pair in held_out["pairs"]]
    with_counterpart = [c for c in held_out_concepts if c in ingested_concepts]
    assert len(with_counterpart) >= 10
    # A couple of held-out questions are deliberately new, so the `new` bucket has a
    # reference answer to score drafts against.
    assert 1 <= len(held_out_concepts) - len(with_counterpart) <= 3


# ---------------------------------------------------------------------------------------------
# Reference documents
# ---------------------------------------------------------------------------------------------


def test_reference_documents_state_their_dates_in_the_first_paragraph(
    out: Path, manifest: dict
) -> None:
    references = [d for d in manifest["documents"] if d["role"] == "reference"]
    assert len(references) == 2
    kinds = {d["expected"]["doc_kind"] for d in references}
    assert kinds <= {"iso_27001", "cyber_essentials_plus"} and len(kinds) == 1
    dates = set()
    for doc in references:
        document = Document(str(out / doc["filename"]))
        body = [p for p in document.paragraphs if p.text.strip() and p.style.name != "Title"]
        first = body[0].text
        date_text = doc["expected"]["effective_date_text"]
        assert f"Certificate issued {date_text}." in first
        assert "ISO/IEC 27001" in first
        assert doc["expected"]["certificate_number"] in first
        dates.add(doc["expected"]["effective_date"])
        assert len(document.paragraphs) >= 8, "enough prose for several chunks"
    assert len(dates) == 2, "different effective dates so supersession fires"


# ---------------------------------------------------------------------------------------------
# Question pack and ground truth
# ---------------------------------------------------------------------------------------------


def test_question_pack_shape(pack_rows: list[tuple]) -> None:
    assert 35 <= len(pack_rows) <= 50
    response_types = [row[5] for row in pack_rows]
    assert response_types.count("Pricing") >= 2
    assert response_types.count("Yes/No") >= 3
    assert response_types.count("Free text") >= 25
    assert {row[6] for row in pack_rows} == {"Yes", "No"}
    for section, number, question, word_limit, _weighting, response_type, _mandatory in pack_rows:
        assert section in content.SECTIONS.values()
        assert number.startswith(section.split(".")[0] + ".")
        assert question.strip()
        if response_type == "Free text":
            assert isinstance(word_limit, int) and word_limit > 0
        else:
            assert word_limit is None
    numbers = [(row[0], row[1]) for row in pack_rows]
    assert len(numbers) == len(set(numbers))


def test_ground_truth_is_consistent_with_the_pack(
    out: Path, ground_truth: dict, pack_rows: list[tuple], manifest: dict
) -> None:
    by_key = {(q["section"], q["number"]): q for q in ground_truth["questions"]}
    assert set(by_key) == {(row[0], row[1]) for row in pack_rows}
    for section, number, question, *_rest in pack_rows:
        assert by_key[(section, number)]["question_text"] == question

    filenames = {doc["filename"] for doc in manifest["documents"]}
    assert ground_truth["question_pack"] == content.QUESTION_PACK_FILENAME
    assert ground_truth["held_out_submission"] in filenames
    assert set(ground_truth["ingested_submissions"]) <= filenames
    assert ground_truth["held_out_submission"] not in ground_truth["ingested_submissions"]

    texts = {name: docx_text(out / name) for name in filenames if name.endswith(".docx")}
    labels = {q["label"] for q in ground_truth["questions"]}
    assert labels == {"counterpart", "new"}
    for question in ground_truth["questions"]:
        if question["label"] == "new":
            assert question["counterparts"] == []
        else:
            assert question["counterparts"]
        for counterpart in question["counterparts"]:
            assert counterpart["submission_filename"] in ground_truth["ingested_submissions"]
            assert counterpart["question_text"] in texts[counterpart["submission_filename"]]
        if question["held_out_answer"]:
            for paragraph in question["held_out_answer"].split("\n\n"):
                assert paragraph in texts[ground_truth["held_out_submission"]]


def test_ground_truth_labels_have_the_expected_mix(ground_truth: dict) -> None:
    questions = ground_truth["questions"]
    new = [q for q in questions if q["label"] == "new"]
    counterpart = [q for q in questions if q["label"] == "counterpart"]
    assert len(new) >= 5, "several questions with no counterpart"
    assert len(counterpart) >= 15
    assert all(q["label"] == "new" for q in questions if q["response_type"] == "pricing")
    with_answer = [q for q in questions if q["held_out_answer"]]
    assert 10 <= len(with_answer) <= 14
    assert any(q["label"] == "new" for q in with_answer)
    # Yes/no questions that a free-text answer nevertheless covers point at that answer.
    yes_no_with_counterpart = [
        q for q in questions if q["response_type"] == "yes_no" and q["label"] == "counterpart"
    ]
    assert yes_no_with_counterpart


# ---------------------------------------------------------------------------------------------
# Claude mode, exercised through the fake provider
# ---------------------------------------------------------------------------------------------


def test_claude_mode_writes_answers_from_the_llm(tmp_path: Path, fake_llm) -> None:  # noqa: ANN001
    from app.config import get_settings

    def scripted(*, user: str, output_model, **_: object):  # noqa: ANN001
        brief = json.loads(user)
        answers = []
        for question in brief["questions"]:
            fixed = question["must_include_verbatim"]
            # Drop the first fixed sentence to check the writer restores it.
            paragraphs = [
                f"Our answer for {question['concept_id']} addresses {brief['buyer']}.",
                " ".join(fixed[1:]) or "We describe our approach in full.",
            ]
            if question.get("base_answer"):
                paragraphs.append("This repeats our earlier response almost word for word.")
            answers.append({"concept_id": question["concept_id"], "paragraphs": paragraphs})
        return output_model.model_validate({"answers": answers})

    fake_llm.register("synthetic_answers", scripted)
    manifest = generate(tmp_path, seed=0, mode="claude")

    assert manifest["mode"] == "claude"
    assert len(fake_llm.calls) == len(content.SUBMISSIONS)
    for call in fake_llm.calls:
        assert call.name == "synthetic_answers"
        assert call.model == get_settings().model_main
        assert "No headings" in call.system
    documents = {d["key"]: d for d in manifest["documents"] if "key" in d}
    for spec in content.SUBMISSIONS:
        doc = documents[spec.key]
        assert doc["writer"] == "claude"
        assert doc["prompt_version"] == "synthetic_answers.v1"
        assert doc["template_fallback_concept_ids"] == []
        answers = answers_by_concept(tmp_path, doc)
        for concept_id, answer in answers.items():
            assert f"Our answer for {concept_id}" in answer
        if spec.key == "westmoor_2024":
            snapshot = content.facts_as_at(spec.submission_date)
            assert snapshot.iso_cert in answers["is_iso27001"], "dropped fact sentence restored"
    # The second submission's brief carried the first's answers for the near-identical pairs.
    second_brief = json.loads(fake_llm.calls[1].user)
    with_base = [q["concept_id"] for q in second_brief["questions"] if "base_answer" in q]
    assert set(with_base) == set(content.SUBMISSIONS[1].near_identical_to)


def test_claude_mode_falls_back_to_templates_for_skipped_questions(
    tmp_path: Path, fake_llm
) -> None:  # noqa: ANN001
    def scripted(*, user: str, output_model, **_: object):  # noqa: ANN001
        brief = json.loads(user)
        first = brief["questions"][0]
        return output_model.model_validate(
            {"answers": [{"concept_id": first["concept_id"], "paragraphs": ["# Heading", "Body."]}]}
        )

    fake_llm.register("synthetic_answers", scripted)
    manifest = generate(tmp_path, seed=0, mode="claude")
    doc = next(d for d in manifest["documents"] if d.get("key") == "westmoor_2024")
    assert set(doc["template_fallback_concept_ids"]) == set(content.SUBMISSIONS[0].concept_ids[1:])
    answers = answers_by_concept(tmp_path, doc)
    assert "# Heading" not in answers[content.SUBMISSIONS[0].concept_ids[0]]
