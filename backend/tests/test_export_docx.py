"""docx export: title, headings, prose, placeholder and review-mode Word comments."""

from __future__ import annotations

import io

from docx import Document
from sqlalchemy.orm import Session

from app.export import PLACEHOLDER, build_export
from app.export.docx import REVIEW_NOTE, source_comment
from tests.test_export_helpers import (
    make_answer,
    make_question,
    make_question_pack,
    make_tender,
    segment,
    traced_segments,
    traced_text,
)


def _open(payload: bytes):
    return Document(io.BytesIO(payload))


def _texts(document) -> list[str]:
    return [paragraph.text for paragraph in document.paragraphs]


def _headings(document) -> list[tuple[str, str]]:
    return [
        (paragraph.style.name, paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith(("Title", "Heading"))
    ]


def test_submission_docx_writes_every_question_in_order_with_prose_or_placeholder(
    db_session: Session,
) -> None:
    tender = make_tender(db_session, name="Community EPR", buyer="NHS Test ICB")
    pack = make_question_pack(db_session, tender)
    second = make_question(
        db_session, tender, pack, order_index=2, section="Section 2: Clinical safety",
        number="2.1", text="Describe your clinical safety process.", status="approved",
    )
    first = make_question(
        db_session, tender, pack, order_index=1, section="Section 1: Company", number="1.1",
        text="Describe your organisation.", status="sme_verified",
    )
    make_question(
        db_session, tender, pack, order_index=3, section="Section 2: Clinical safety",
        number="2.2", text="Name your clinical safety officer.",
    )
    make_answer(db_session, second, text="We follow DCB0129.\n\nA hazard log is maintained.")
    make_answer(db_session, first, text="Verified but not yet approved.")

    document = _open(build_export(db_session, tender, format="docx", mode="submission"))

    assert document.core_properties.title == "Community EPR – NHS Test ICB"
    assert _headings(document) == [
        ("Title", "Community EPR – NHS Test ICB"),
        ("Heading 1", "Section 1: Company"),
        ("Heading 2", "Question 1.1"),
        ("Heading 1", "Section 2: Clinical safety"),
        ("Heading 2", "Question 2.1"),
        ("Heading 2", "Question 2.2"),
    ], "one heading per question in order_index; a section heading when the section changes"

    texts = _texts(document)
    # Question text precedes its answer; sme_verified is not approved, so it gets the placeholder.
    index_first = texts.index("Describe your organisation.")
    assert texts[index_first + 1] == PLACEHOLDER
    # The approved answer's blank line becomes a paragraph break.
    index_second = texts.index("Describe your clinical safety process.")
    assert texts[index_second + 1 : index_second + 3] == [
        "We follow DCB0129.",
        "A hazard log is maintained.",
    ]
    index_third = texts.index("Name your clinical safety officer.")
    assert texts[index_third + 1] == PLACEHOLDER
    assert texts.count(PLACEHOLDER) == 2
    assert "Verified but not yet approved." not in texts
    assert REVIEW_NOTE not in texts
    assert len(list(document.comments)) == 0, "submission export is clean prose"


def test_title_without_buyer_is_the_tender_name(db_session: Session) -> None:
    tender = make_tender(db_session, name="Solo tender", buyer=None)
    document = _open(build_export(db_session, tender, format="docx", mode="submission"))
    assert _headings(document) == [("Title", "Solo tender")]


def test_review_docx_attaches_sources_as_word_comments(db_session: Session) -> None:
    tender = make_tender(db_session)
    pack = make_question_pack(db_session, tender)
    drafted = make_question(
        db_session, tender, pack, order_index=1, section="Clinical safety", number="3.2",
        text="Describe your clinical safety arrangements.", status="ai_draft",
        needs_review=True,
    )
    empty = make_question(
        db_session, tender, pack, order_index=2, section="Clinical safety", number="3.3",
        text="Attach your safety case.",
    )
    make_answer(db_session, drafted, text=traced_text(), segments=traced_segments())

    # needs_review does not block review mode.
    document = _open(build_export(db_session, tender, format="docx", mode="review"))

    texts = _texts(document)
    assert REVIEW_NOTE in texts
    index = texts.index("Describe your clinical safety arrangements.")
    assert texts[index + 1] == (
        "We hold a clinical safety case for the product. It was last reviewed in January 2026. "
        "In addition,"
    ), "sentences joined with a single space within a paragraph"
    assert texts[index + 2] == (
        "Our named clinical safety officer chairs the hazard log review. Every release passes a "
        "full regression suite. The hazard log is shared with the buyer quarterly."
    )
    assert texts[texts.index("Attach your safety case.") + 1] == PLACEHOLDER
    assert empty.id  # the placeholder question exists and was written

    comments = list(document.comments)
    # Six segments; the connective one with no sources gets no comment.
    assert len(comments) == 5
    by_text = [comment.text for comment in comments]

    supported = by_text[0]
    assert supported.startswith("Support: Supported")
    assert "Past submission 2025.docx (past_submission)" in supported
    assert "Effective 2025-03-01" in supported
    assert "page 4" in supported and "3.2 Clinical safety" in supported
    assert "“clinical safety case is maintained for the product”" in supported

    two_sources = by_text[1]
    assert "1. Source: Past submission 2025.docx" in two_sources
    assert "2. Fact: ISO 27001 certificate.pdf (reference, iso_27001)" in two_sources
    assert "table 1" in two_sources and "row Sheet1!5" in two_sources

    attested = by_text[2]
    assert attested.startswith("Support: Human-authored")
    assert "Attested by Sam Reviewer at 2026-09-28T10:12:00Z: Confirmed" in attested

    unsupported = by_text[3]
    assert unsupported == "Support: Unsupported\nNo source located for this sentence."

    disputed = by_text[4]
    assert "Disputed by Sam Reviewer at 2026-09-28T11:00:00Z: We share it monthly" in disputed
    assert "span not located in this section" in disputed

    # Every comment is anchored to a run in the body.
    body_xml = document.element.body.xml
    assert body_xml.count("commentRangeStart") == 5
    assert body_xml.count("commentReference") == 5


def test_review_docx_falls_back_to_prose_when_a_version_has_no_segments(
    db_session: Session,
) -> None:
    tender = make_tender(db_session)
    pack = make_question_pack(db_session, tender)
    question = make_question(
        db_session, tender, pack, order_index=1, section="Commercial", number="C1",
        text="Provide your pricing schedule.", response_type="pricing", status="writer_edited",
    )
    make_answer(
        db_session, question, text="Priced per user.\n\nDiscounts apply.", author_type="user"
    )

    document = _open(build_export(db_session, tender, format="docx", mode="review"))
    texts = _texts(document)
    index = texts.index("Provide your pricing schedule.")
    assert texts[index + 1 : index + 3] == ["Priced per user.", "Discounts apply."]
    assert len(list(document.comments)) == 0


def test_source_comment_omits_connective_without_sources() -> None:
    connective = segment(0, 0, "However,", kind="connective", support_status="connective")
    assert source_comment(connective) is None
    assert source_comment(segment(0, 0, "Claim.", support_status="unsupported")) == (
        "Support: Unsupported\nNo source located for this sentence."
    )
