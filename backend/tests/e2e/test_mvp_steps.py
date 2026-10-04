"""The six MVP steps of CLAUDE.md, end to end, offline, on the synthetic data in ``eval/data``.

Everything goes through the real API (``app_client``) and the real job handlers, dispatched
by the worker's ``run_once`` in a loop instead of a background process. The only doubles are
the heuristic language model (``tests.fakes.heuristic_llm``) and the hash-based embeddings.
Step 6 (the evaluation report) belongs to the harness and is not exercised here.

The test is one function so the state flows as it would for a user; each step is a helper
whose name is the step it proves, and every assertion says which step it belongs to.
"""

from __future__ import annotations

import io
import json
import uuid
from datetime import UTC, datetime, timedelta
from typing import Any

import httpx
from docx import Document as read_docx
from sqlalchemy.orm import Session

from app.db.models import Fact
from app.export import PLACEHOLDER
from app.generate import runner
from app.review.invalidation import run_expiry_sweep
from tests.e2e.conftest import EvalData
from tests.fakes.heuristic_llm import HeuristicFakeLLM

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ACTOR = "test user"
STORED_STATUSES = {"supported", "weak", "unsupported", "human_authored", "connective"}


# --- Helpers ----------------------------------------------------------------------------------


async def get(client: httpx.AsyncClient, url: str, **params: Any) -> Any:
    response = await client.get(url, params=params or None)
    assert response.status_code == 200, f"GET {url}: {response.status_code} {response.text}"
    return response.json()


async def post(client: httpx.AsyncClient, url: str, body: Any = None, *, expect: int) -> Any:
    response = await client.post(url, json=body)
    assert response.status_code == expect, f"POST {url}: {response.status_code} {response.text}"
    return response.json() if response.content else None


async def patch(client: httpx.AsyncClient, url: str, body: Any, *, expect: int = 200) -> Any:
    response = await client.patch(url, json=body)
    assert response.status_code == expect, f"PATCH {url}: {response.status_code} {response.text}"
    return response.json()


async def stream(client: httpx.AsyncClient, url: str, body: dict[str, Any]) -> list[dict]:
    """Consume an NDJSON stream as the front end would, one JSON object per line."""
    async with client.stream("POST", url, json=body) as response:
        if response.status_code != 200:
            await response.aread()
            raise AssertionError(f"POST {url}: {response.status_code} {response.text}")
        assert response.headers["content-type"].startswith("application/x-ndjson")
        return [json.loads(line) async for line in response.aiter_lines() if line.strip()]


def substantive(answer: dict[str, Any]) -> list[dict[str, Any]]:
    return [segment for segment in answer["segments"] if segment["kind"] == "substantive"]


def has_source(segment: dict[str, Any], source_type: str) -> bool:
    return any(source["source_type"] == source_type for source in segment["sources"])


def text_from_segments(segments: list[dict[str, Any]]) -> str:
    """The plan's derivation: sentences joined by a space, paragraphs by a blank line."""
    paragraphs: dict[int, list[str]] = {}
    for segment in segments:
        paragraphs.setdefault(int(segment["paragraph"]), []).append(segment["text"])
    return "\n\n".join(" ".join(paragraphs[key]) for key in sorted(paragraphs))


def transitions(question: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {entry["to"]: entry for entry in question["allowed_transitions"]}


# --- Step 1: the library ----------------------------------------------------------------------


async def step_1_library(
    client: httpx.AsyncClient, eval_data: EvalData, run_jobs, db_session: Session
) -> dict[str, Any]:
    """Upload three past submissions and two reference documents, ingest them to ``ready``,
    confirm each classification, and see the older certificate superseded."""
    uploaded: dict[str, dict[str, Any]] = {}
    roles = [("past_submission", None), ("held_out_submission", None), ("reference", None)]
    for document in eval_data.manifest["documents"]:
        if document["role"] not in {role for role, _ in roles}:
            continue
        path = eval_data.path(document["filename"])
        response = await client.post(
            "/documents", files={"file": (document["filename"], path.read_bytes(), DOCX)}
        )
        assert response.status_code == 202, response.text
        body = response.json()
        assert body["job_id"] and body["ingest_status"] == "queued"
        uploaded[document["filename"]] = body

    ran = run_jobs()
    assert ran == len(uploaded), f"step 1: expected {len(uploaded)} ingest jobs, ran {ran}"

    listing = {row["filename"]: row for row in await get(client, "/documents")}
    for document in eval_data.manifest["documents"]:
        filename = document["filename"]
        if filename not in uploaded:
            continue
        row = listing[filename]
        expected = document["expected"]
        assert row["ingest_status"] == "ready", f"step 1: {filename} is {row['ingest_status']}"
        assert row["doc_type"] == expected["doc_type"], filename
        assert row["classification_confirmed"] is False, "confirmation is the user's step"
        assert row["effective_date"] == expected["effective_date"], filename
        assert row["effective_date_source"] == "extracted", filename
        if expected["doc_type"] == "past_submission":
            assert row["buyer"] == expected["buyer"]
            assert row["submission_date"] == expected["submission_date"]
            assert row["item_counts"]["qa_pair"] == expected["pair_count"], (
                f"step 1: {filename} extracted {row['item_counts']} pairs"
            )
            assert row["fact_count"] >= 3, f"step 1: {filename} has {row['fact_count']} facts"
        else:
            assert row["doc_kind"] == expected["doc_kind"]
            assert row["item_counts"]["chunk"] >= 1
            assert row["fact_count"] >= 1
        job = await get(client, f"/jobs/{uploaded[filename]['job_id']}")
        assert job["status"] == "done" and job["done"] == job["total"] == 5, job

    # Every extracted item was sliced from, or found in, its section: the fidelity target.
    items = await get(client, "/library/items")
    pairs = [item for item in items if item["item_type"] == "qa_pair"]
    assert pairs and all(item["text_verified"] for item in pairs), "step 1: unverified pairs"
    unpaired = await get(
        client,
        f"/documents/{listing[eval_data.document('past_submission')['filename']]['id']}/unpaired",
    )
    assert unpaired["items"] == [] and unpaired["fragments"] == []

    # Sections of one document open in the source pane's shape.
    first = eval_data.document("past_submission")
    sections = await get(client, f"/documents/{listing[first['filename']]['id']}/sections")
    assert len(sections) >= 20 and [s["order_index"] for s in sections] == list(
        range(len(sections))
    )
    assert all(section["text"].strip() for section in sections)
    assert any(section["cell_ref"] for section in sections), "table rows carry cell_ref"

    # Confirm the submissions, then the older certificate, then the newer one.
    for document in eval_data.manifest["documents"]:
        if document["role"] in ("past_submission", "held_out_submission"):
            confirmed = await post(
                client, f"/documents/{listing[document['filename']]['id']}/confirm", expect=200
            )
            assert confirmed["classification_confirmed"] is True
    older = listing[eval_data.document("reference", "iso_2024")["filename"]]
    newer = listing[eval_data.document("reference", "iso_2025")["filename"]]
    confirmed_older = await post(client, f"/documents/{older['id']}/confirm", expect=200)
    assert confirmed_older["superseded_by"] is None, "one confirmed certificate: nothing to pair"
    assert await get(client, "/library/supersession-decisions") == []
    confirmed_newer = await post(client, f"/documents/{newer['id']}/confirm", expect=200)
    assert confirmed_newer["superseded_by"] is None
    older_now = await get(client, f"/documents/{older['id']}")
    assert older_now["superseded_by"] == newer["id"], "step 1: the 2024 certificate is superseded"
    assert await get(client, "/library/supersession-decisions") == [], "decided automatically"
    items = await get(client, "/library/items")
    for item in items:
        if item["document_id"] == older["id"]:
            assert item["excluded_from_retrieval"] is True
        else:
            assert item["excluded_from_retrieval"] is False

    # Deduplication found the near-identical answers and facts carry keys and a chain.
    clustered = [item for item in items if item["canonical_id"] is not None]
    assert clustered, "step 1: no deduplication cluster formed"
    facts = await get(client, "/library/facts")
    keyed = {fact["fact_key"] for fact in facts if fact["fact_key"]}
    assert keyed, "step 1: no fact carries a key"
    kinds = {fact["fact_kind"] for fact in facts}
    assert {"iso_27001", "cyber_essentials_plus", "dspt_status"} <= kinds, kinds
    superseded_facts = [fact for fact in facts if fact["superseded_by"] is not None]
    assert superseded_facts, "step 1: the renewed certificates supersede the earlier facts"
    assert all(fact["is_current"] is False for fact in superseded_facts)
    assert any(fact["is_current"] for fact in facts if fact["fact_kind"] == "iso_27001")

    return {"documents": listing, "older_reference": older, "newer_reference": newer}


# --- Step 2: a tender and its question pack ---------------------------------------------------


async def step_2_tender(client: httpx.AsyncClient, eval_data: EvalData, run_jobs) -> dict[str, Any]:
    pack = eval_data.document("question_pack")
    tender = await post(
        client,
        "/tenders",
        {
            "name": "Community clinical communication platform",
            "buyer": pack["expected"]["buyer"],
            "deadline": (datetime.now(UTC) + timedelta(days=21)).isoformat(),
        },
        expect=201,
    )
    response = await client.post(
        f"/tenders/{tender['id']}/documents",
        data={"tender_doc_kind": "question_pack"},
        files={"file": (pack["filename"], eval_data.path(pack["filename"]).read_bytes(), XLSX)},
    )
    assert response.status_code == 202, response.text
    upload = response.json()
    extract_job_id = upload["job_id"]

    ran = run_jobs()
    assert ran == 2, f"step 2: expected extract then triage, ran {ran} job(s)"
    extract_job = await get(client, f"/jobs/{extract_job_id}")
    assert extract_job["status"] == "done", extract_job["error"]
    assert extract_job["next_job_id"], "extraction chains the triage job"
    assert all(result["outcome"] == "extracted" for result in extract_job["results"])
    triage_job = await get(client, f"/jobs/{extract_job['next_job_id']}")
    assert triage_job["status"] == "done", triage_job["error"]
    assert triage_job["total"] == triage_job["done"] == pack["expected"]["question_count"]

    detail = await get(client, f"/tenders/{tender['id']}")
    assert detail["extract_job_id"] == extract_job_id
    assert detail["triage_job_id"] == extract_job["next_job_id"]
    assert detail["questions_total"] == pack["expected"]["question_count"]
    assert detail["documents"][0]["ingest_status"] == "ready"

    questions = await get(client, f"/tenders/{tender['id']}/questions")
    assert len(questions) == pack["expected"]["question_count"]
    assert [q["order_index"] for q in questions] == list(range(len(questions)))
    assert all(q["thread_id"] for q in questions), "one thread per question"
    assert all(q["status"] == "not_started" and q["current_answer"] is None for q in questions)
    assert all(q["coverage"] != "unknown" for q in questions), "step 2: a question left unknown"
    assert sum(1 for q in questions if q["topics"]) >= len(questions) * 0.9

    by_number = {(q["section"], q["number"]): q for q in questions}
    truth = {(t["section"], t["number"]): t for t in eval_data.ground_truth["questions"]}
    assert set(by_number) == set(truth), "every pack row became a card"
    distribution = {label: 0 for label in ("covered", "partial", "new")}
    for question in questions:
        distribution[question["coverage"]] += 1
    counterpart_hit = sum(
        1
        for key, entry in truth.items()
        if entry["label"] == "counterpart" and by_number[key]["coverage"] in ("covered", "partial")
    )
    counterpart_total = sum(1 for entry in truth.values() if entry["label"] == "counterpart")
    new_hit = [
        key
        for key, entry in truth.items()
        if entry["label"] == "new" and by_number[key]["coverage"] == "new"
    ]
    print(
        f"step 2 coverage: {distribution}; counterparts covered or partial "
        f"{counterpart_hit}/{counterpart_total}; ground-truth new judged new: {len(new_hit)}"
    )
    assert distribution["covered"] >= 8, distribution
    assert distribution["new"] >= 1, distribution
    assert new_hit, "at least one question the manifest marks new is triaged new"
    assert counterpart_hit >= counterpart_total * 0.75
    for question in questions:
        detail = question["coverage_detail"]
        assert detail["label_source"] in ("floor", "llm")
        assert "best_vec" in detail and "candidates" in detail
        assert detail["label_source"] == "llm" or question["coverage"] == "new"
    return {"tender": tender, "questions": questions, "by_number": by_number}


# --- Step 3: a streamed draft with a verified trace -------------------------------------------


async def step_3_draft(
    client: httpx.AsyncClient,
    db_session: Session,
    llm: HeuristicFakeLLM,
    questions: list[dict[str, Any]],
) -> dict[str, Any]:
    covered = [
        q for q in questions if q["coverage"] == "covered" and q["response_type"] == "free_text"
    ]
    assert covered, "step 3 needs a covered free-text question"
    question = next((q for q in covered if q["number"] == "3.1"), covered[0])

    events = await stream(client, f"/questions/{question['id']}/draft", {})
    await runner.wait_for(runner.question_key(question["id"]))
    db_session.expire_all()

    types = [event["type"] for event in events]
    offer: dict[str, Any] | None = None
    if types[0] == "verbatim":
        offer = events[0]
        assert offer["source_item_id"] and offer["segments"], "the verbatim offer is complete"
        assert all(s["support_status"] == "supported" for s in offer["segments"])
        types = types[1:]
        events_after_offer = events[1:]
    else:
        events_after_offer = events
    segment_count = types.count("segment")
    assert segment_count >= 3, types
    assert types == (
        ["segment"] * segment_count
        + ["gaps", "fact_checklist"]
        + ["support"] * segment_count
        + ["done"]
    ), f"step 3: event order {types}"

    segments = [event for event in events_after_offer if event["type"] == "segment"]
    supports = [event for event in events_after_offer if event["type"] == "support"]
    done = events[-1]
    assert [s["index"] for s in segments] == list(range(segment_count)), "wire indices"
    assert all(s["support_status"] == "pending" for s in segments), "pending on the wire only"
    assert sorted(s["index"] for s in supports) == list(range(segment_count))
    assert [s["index"] for s in done["segments"]] == list(range(segment_count)), (
        "step 3: segment indices equal the persisted indices"
    )
    assert [s["text"] for s in done["segments"]] == [s["text"] for s in segments]
    assert all(s["support_status"] in STORED_STATUSES for s in done["segments"])
    assert llm.synthesis_runs[-1].merged_segment
    assert segment_count > llm.synthesis_runs[-1].model_segments, (
        "conformance split the two-sentence model segment"
    )

    # Support events carry offsets that resolve through GET /sections to the quoted text.
    located = 0
    fact_sources = 0
    for support in supports:
        assert support["support_status"] in STORED_STATUSES
        for source in support["sources"]:
            if source["source_type"] == "human_attestation":
                continue
            if source["source_type"] == "fact":
                fact_sources += 1
            locator = source["locator"]
            if locator["start"] is None or locator["end"] is None:
                continue
            located += 1
            body = await get(
                client,
                f"/sections/{locator['section_id']}",
                start=locator["start"],
                end=locator["end"],
            )
            assert body["highlight"] == {"start": locator["start"], "end": locator["end"]}
            assert body["section"]["document_id"] == locator["document_id"]
            assert body["section"]["text"][locator["start"] : locator["end"]] == source["quote"]
            assert body["document"]["filename"] == source["document_title"]
    assert located >= segment_count - 1, "step 3: every substantive sentence located its span"
    assert fact_sources >= 1, "step 3: the draft cites at least one fact"
    supported = [s for s in done["segments"] if s["support_status"] == "supported"]
    assert supported, "step 3: nothing verified as supported"

    # done carries the persisted answer.
    answers = await get(client, f"/questions/{question['id']}/answers")
    assert len(answers) == 1 and answers[0]["id"] == done["id"]
    assert answers[0]["is_current"] and answers[0]["version"] == 1
    assert answers[0]["segments"] == done["segments"]
    assert answers[0]["support_summary"] == done["support_summary"]
    assert done["model"] and done["prompt_version"] == "synthesis.v2"
    assert done["word_count"] == len(done["text"].split())
    if offer is not None:
        # The pack question repeats a past question word for word: the offer is persisted by
        # item id only and rebuilt on read.
        assert done["verbatim_offer_item_id"] == offer["source_item_id"]
        assert done["verbatim"]["source_item_id"] == offer["source_item_id"]
        assert [s["text"] for s in done["verbatim"]["segments"]] == [
            s["text"] for s in offer["segments"]
        ]
    else:
        assert done["verbatim_offer_item_id"] is None and done["verbatim"] is None
    print(f"step 3: verbatim offer {'made' if offer else 'not made'}; {segment_count} segments")
    assert done["fact_checklist"], "the fact checklist names the facts used"
    assert all(
        entry["status"] in ("current", "superseded", "expired", "unverified")
        for entry in done["fact_checklist"]
    )
    detail = await get(client, f"/questions/{question['id']}")
    assert detail["status"] == "ai_draft"
    assert detail["draft_in_progress"] is False
    assert detail["current_answer"]["id"] == done["id"]
    assert transitions(detail)["ai_draft"]["blockers"] == [
        {"kind": "system_only", "index": None, "gap": None}
    ]
    return {"question": detail, "answer": answers[0], "events": events}


# --- Step 4: the review pipeline --------------------------------------------------------------


async def step_4_review(
    client: httpx.AsyncClient,
    db_session: Session,
    question: dict[str, Any],
    evidence_document_id: str,
) -> dict[str, Any]:
    qid = question["id"]
    thread_id = question["thread_id"]

    # The assistant thread: an instruction, then save the reply as the answer.
    events = await stream(
        client, f"/threads/{thread_id}/messages", {"content": "Make this shorter."}
    )
    await runner.wait_for(runner.thread_key(thread_id))
    db_session.expire_all()
    assert events[-1]["type"] == "done" and events[-1]["role"] == "assistant"
    reply = events[-1]
    thread = await get(client, f"/threads/{thread_id}")
    assert [m["role"] for m in thread["messages"]] == ["user", "assistant"]
    assert thread["messages"][1]["id"] == reply["id"]
    assert thread["messages"][1]["segments"], "a message is a list of segments"
    saved = await post(client, f"/messages/{reply['id']}/save-as-answer", {}, expect=201)
    v2 = saved["answer"]
    assert v2["version"] == 2 and v2["is_current"] and v2["author_type"] == "ai"
    assert [s["text"] for s in v2["segments"]] == [s["text"] for s in reply["segments"]]
    assert saved["question"]["status"] == "ai_draft"
    assert len(v2["segments"]) < len(question["current_answer"]["segments"]), "shortened"
    thread = await get(client, f"/threads/{thread_id}")
    assert thread["messages"][1]["answer_id"] == v2["id"]

    # A human edit with unchanged text keeps every status and lands at writer_edited.
    edited = await post(
        client,
        f"/questions/{qid}/answers",
        {"text": v2["text"], "base_version_id": v2["id"]},
        expect=201,
    )
    v3 = edited["answer"]
    assert v3["version"] == 3 and v3["author_type"] == "user" and v3["author_name"] == ACTOR
    assert [(s["text"], s["support_status"]) for s in v3["segments"]] == [
        (s["text"], s["support_status"]) for s in v2["segments"]
    ]
    assert [s["sources"] for s in v3["segments"]] == [s["sources"] for s in v2["segments"]]
    assert edited["question"]["status"] == "writer_edited"
    assert v3["gaps"] == v2["gaps"] and v3["fact_checklist"] == v2["fact_checklist"]
    stale = await client.post(
        f"/questions/{qid}/answers", json={"text": v2["text"], "base_version_id": v2["id"]}
    )
    assert stale.status_code == 409 and stale.json()["current_answer"]["id"] == v3["id"]

    # A changed sentence drops to weak and is re-verified; an added sentence is human-authored.
    target = next(
        s
        for s in substantive(v3)
        if s["support_status"] == "supported" and has_source(s, "knowledge_item")
    )
    words = target["text"].split()
    changed_text = " ".join([words[0], "clearly", *words[1:]])
    added_text = "We will supply supporting evidence for this answer on request."
    new_segments = [dict(s) for s in v3["segments"]]
    for s in new_segments:
        if s["index"] == target["index"]:
            s["text"] = changed_text
    last_paragraph = max(int(s["paragraph"]) for s in new_segments)
    new_segments.append({"text": added_text, "paragraph": last_paragraph})
    edited = await post(
        client,
        f"/questions/{qid}/answers",
        {"text": text_from_segments(new_segments), "base_version_id": v3["id"]},
        expect=201,
    )
    v4 = edited["answer"]
    changed = next(s for s in v4["segments"] if s["text"] == changed_text)
    assert changed["support_status"] == "supported", "re-verified after dropping to weak"
    assert has_source(changed, "knowledge_item"), "kept its document sources"
    added = next(s for s in v4["segments"] if s["text"] == added_text)
    assert added["support_status"] == "human_authored" and added["sources"] == []
    for s in v4["segments"]:
        if s["text"] not in (changed_text, added_text):
            before = next(b for b in v3["segments"] if b["text"] == s["text"])
            assert s["support_status"] == before["support_status"]
    events_now = await get(client, f"/questions/{qid}/events")
    created_v4 = next(
        e
        for e in events_now
        if e["event_type"] == "answer_created" and e["payload"].get("answer_id") == v4["id"]
    )
    assert created_v4["payload"]["reverified"] == 1, "exactly the changed sentence went weak"

    # Attest the human-authored sentence.
    attested = await post(
        client,
        f"/answers/{v4['id']}/segments/{added['index']}/attest",
        {"note": "Confirmed with the bid team."},
        expect=200,
    )
    segment = attested["answer"]["segments"][added["index"]]
    assert segment["support_status"] == "supported"
    assert segment["sources"][-1]["source_type"] == "human_attestation"
    assert segment["sources"][-1]["attested_by"] == ACTOR
    refused = await client.post(
        f"/answers/{v4['id']}/segments/{added['index']}/attest", json={"note": "again"}
    )
    assert refused.status_code == 409, "a supported sentence cannot be attested"

    # Dispute a supported sentence, then rewrite it and attest the rewrite. Prefer a sentence
    # that cites no fact, so a fact-cited sentence survives for the expiry path later.
    disputable = [
        s
        for s in substantive(attested["answer"])
        if s["support_status"] == "supported"
        and has_source(s, "knowledge_item")
        and s["text"] != changed_text
    ]
    assert disputable, "step 4 needs a supported sentence to dispute"
    disputed_target = next((s for s in disputable if not has_source(s, "fact")), disputable[-1])
    disputed = await post(
        client,
        f"/answers/{v4['id']}/segments/{disputed_target['index']}/dispute",
        {"note": "The audit date needs checking against the latest report."},
        expect=200,
    )
    segment = disputed["answer"]["segments"][disputed_target["index"]]
    assert segment["support_status"] == "unsupported" and segment["dispute"]["disputed_by"] == ACTOR
    assert segment["sources"], "a disputed sentence keeps its sources for context"
    assert disputed["question"]["needs_review"] is True
    blockers = transitions(disputed["question"])["sme_verified"]["blockers"]
    assert {"kind": "needs_review", "index": None, "gap": None} in blockers
    assert any(b["kind"] == "segment" and b["index"] == disputed_target["index"] for b in blockers)

    rewritten_text = (
        "Our bid manager has re-checked this statement against the current certificate on file."
    )
    rewrite_segments = [dict(s) for s in disputed["answer"]["segments"]]
    for s in rewrite_segments:
        if s["index"] == disputed_target["index"]:
            s["text"] = rewritten_text
    edited = await post(
        client,
        f"/questions/{qid}/answers",
        {"text": text_from_segments(rewrite_segments), "base_version_id": v4["id"]},
        expect=201,
    )
    v5 = edited["answer"]
    rewritten = next(s for s in v5["segments"] if s["text"] == rewritten_text)
    assert rewritten["support_status"] == "human_authored" and rewritten["dispute"] is None
    kept_attested = next(s for s in v5["segments"] if s["text"] == added_text)
    assert kept_attested["support_status"] == "supported", "an unchanged attestation survives"
    await post(
        client,
        f"/answers/{v5['id']}/segments/{rewritten['index']}/attest",
        {"note": "Rewritten after the dispute."},
        expect=200,
    )
    # Anything else still short of supported (a sentence resting on a superseded fact, for
    # example) is attested so the SME gate can be shown passing.
    current = (await get(client, f"/questions/{qid}"))["current_answer"]
    for s in substantive(current):
        if s["support_status"] != "supported":
            await post(
                client,
                f"/answers/{current['id']}/segments/{s['index']}/attest",
                {"note": "Checked against the current certificate."},
                expect=200,
            )
    detail = await get(client, f"/questions/{qid}")
    assert detail["needs_review"] is False, "the recompute cleared the flag"
    assert all(s["support_status"] == "supported" for s in substantive(detail["current_answer"]))
    assert detail["current_answer"]["support_summary"]["score"] == 1.0

    # Gaps must be acknowledged before approval.
    gaps = detail["current_answer"]["gaps"]
    assert gaps, "step 4 needs a gap to acknowledge"
    assert transitions(detail)["sme_verified"]["allowed"] is True
    assert transitions(detail)["approved"]["allowed"] is False
    assert {b["gap"] for b in transitions(detail)["approved"]["blockers"]} == set(gaps)
    for gap in gaps:
        detail = await post(
            client,
            f"/questions/{qid}/gaps/acknowledge",
            {"gap": gap, "note": "Not required by this buyer."},
            expect=200,
        )
    assert {ack["gap"] for ack in detail["gap_acknowledgements"]} == set(gaps)
    unknown_gap = await client.post(
        f"/questions/{qid}/gaps/acknowledge", json={"gap": "Not a gap", "note": None}
    )
    assert unknown_gap.status_code == 422

    # Compliance class, evidence, comment, assignment.
    detail = await patch(client, f"/questions/{qid}", {"compliance_class": "A"})
    assert detail["compliance_class"] == "A"
    detail = await post(
        client,
        f"/questions/{qid}/evidence",
        {"document_id": evidence_document_id, "note": "Current ISO 27001 certificate"},
        expect=201,
    )
    assert detail["evidence"][0]["document_id"] == evidence_document_id
    comment = await post(
        client,
        f"/questions/{qid}/comments",
        {"text": "Checked against the 2025 certificate."},
        expect=201,
    )
    assert comment["author"] == ACTOR
    detail = await patch(client, f"/questions/{qid}", {"assignee": "Priya Nair"})
    assert detail["assignee"] == "Priya Nair"
    assert len(detail["comments"]) == 1

    # Through the review states to approved, consulting allowed_transitions first.
    assert transitions(detail)["sme_verified"]["allowed"] is True
    detail = await patch(client, f"/questions/{qid}", {"status": "sme_verified"})
    assert detail["status"] == "sme_verified"
    assert transitions(detail)["approved"]["allowed"] is True
    detail = await patch(client, f"/questions/{qid}", {"status": "approved"})
    assert detail["status"] == "approved"
    blocked = await client.patch(f"/questions/{qid}", json={"status": "ai_draft"})
    assert blocked.status_code == 409 and blocked.json()["blockers"] == [{"kind": "system_only"}]

    # The event log is complete.
    events_now = await get(client, f"/questions/{qid}/events")
    kinds = [e["event_type"] for e in events_now]
    expected = {
        "answer_created",
        "answer_saved_from_chat",
        "status_changed",
        "attested",
        "disputed",
        "gap_acknowledged",
        "compliance_class_set",
        "evidence_added",
        "comment_added",
        "assigned",
    }
    assert expected <= set(kinds), sorted(set(kinds))
    assert all(e["actor"] == ACTOR for e in events_now), "every step names the person"
    status_changes = [e["payload"]["to"] for e in events_now if e["event_type"] == "status_changed"]
    assert status_changes[0] == "approved" and status_changes[1] == "sme_verified"
    assert "writer_edited" in status_changes and "ai_draft" in status_changes
    versions = await get(client, f"/questions/{qid}/answers")
    assert [v["version"] for v in versions] == [1, 2, 3, 4, 5]
    assert [v["is_current"] for v in versions] == [False, False, False, False, True]
    return {"question": detail, "answer": versions[-1]}


# --- Step 5: table view and export ------------------------------------------------------------


async def step_5_export(
    client: httpx.AsyncClient, tender_id: str, approved_question: dict[str, Any]
) -> dict[str, Any]:
    rows = await get(client, f"/tenders/{tender_id}/questions")
    not_yet_approved = [row for row in rows if row["status"] != "approved"]
    assert len(not_yet_approved) == len(rows) - 1
    approved_row = next(row for row in rows if row["id"] == approved_question["id"])
    assert approved_row["current_answer"]["version"] == 5
    assert approved_row["current_answer"]["support_summary"]["score"] == 1.0
    assert approved_row["assignee"] == "Priya Nair" and approved_row["compliance_class"] == "A"

    response = await client.get(
        f"/tenders/{tender_id}/export", params={"format": "docx", "mode": "submission"}
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(DOCX)
    assert "attachment" in response.headers["content-disposition"]
    document = read_docx(io.BytesIO(response.content))
    paragraphs = [(p.style.name, p.text) for p in document.paragraphs]
    headings = [text for style, text in paragraphs if style == "Heading 2"]
    assert headings == [f"Question {row['number']}" for row in rows], "the buyer's order"
    placeholders = [text for _, text in paragraphs if text == PLACEHOLDER]
    assert len(placeholders) == len(rows) - 1, "every unapproved question gets the placeholder"
    texts = [text for _, text in paragraphs]
    start = texts.index(f"Question {approved_row['number']}")
    assert texts[start + 1] == approved_row["text"]
    answer_paragraphs = [p for p in approved_question["current_answer"]["text"].split("\n\n") if p]
    assert texts[start + 2 : start + 2 + len(answer_paragraphs)] == answer_paragraphs, (
        "step 5: the approved answer is written in full"
    )
    assert PLACEHOLDER not in texts[start : start + 2 + len(answer_paragraphs)]
    tender = await get(client, f"/tenders/{tender_id}")
    assert tender["questions_approved"] == 1
    assert tender["words_approved"] == approved_row["current_answer"]["word_count"]
    assert tender["needs_review_count"] == 0
    return {"rows": rows}


# --- Fact expiry --------------------------------------------------------------------------------


async def fact_expiry_path(
    client: httpx.AsyncClient, db_session: Session, tender_id: str, question_id: str
) -> None:
    detail = await get(client, f"/questions/{question_id}")
    answer = detail["current_answer"]
    cited = next(
        (source["source_id"], segment)
        for segment in substantive(answer)
        if segment["support_status"] == "supported"
        for source in segment["sources"]
        if source["source_type"] == "fact"
    )
    fact_id, cited_segment = cited
    fact = db_session.get(Fact, uuid.UUID(fact_id))
    assert fact is not None
    fact.expires_on = datetime.now(UTC).date() - timedelta(days=1)
    db_session.flush()

    expired = run_expiry_sweep(db_session)
    db_session.commit()
    assert expired == 1
    assert run_expiry_sweep(db_session) == 0, "a second sweep is a no-op"
    db_session.commit()
    db_session.expire_all()

    detail = await get(client, f"/questions/{question_id}")
    assert detail["needs_review"] is True
    answer = detail["current_answer"]
    weak = [s for s in answer["segments"] if s["support_status"] == "weak"]
    assert cited_segment["index"] in {s["index"] for s in weak}
    assert all(has_source(s, "fact") for s in weak)
    assert any(
        entry["fact_id"] == fact_id and entry["status"] == "expired"
        for entry in answer["fact_checklist"]
    )
    events = await get(client, f"/questions/{question_id}/events")
    needs_review_events = [e for e in events if e["event_type"] == "answer_needs_review"]
    assert len(needs_review_events) == 1 and needs_review_events[0]["actor"] == "system"
    assert needs_review_events[0]["payload"]["reason"] == "expired"
    assert transitions(detail)["sme_verified"]["allowed"] is False
    assert detail["status"] == "approved", "status never changes by itself"
    assert (await get(client, f"/tenders/{tender_id}"))["needs_review_count"] == 1

    response = await client.get(
        f"/tenders/{tender_id}/export", params={"format": "docx", "mode": "submission"}
    )
    assert response.status_code == 409, response.text
    blocked = response.json()
    assert blocked["code"] == "needs_review"
    assert isinstance(blocked["detail"], str) and blocked["detail"]
    assert blocked["question_ids"] == [question_id]

    for segment in weak:
        await post(
            client,
            f"/answers/{answer['id']}/segments/{segment['index']}/attest",
            {"note": "Renewal certificate received; figure confirmed."},
            expect=200,
        )
    detail = await get(client, f"/questions/{question_id}")
    assert detail["needs_review"] is False
    response = await client.get(
        f"/tenders/{tender_id}/export", params={"format": "docx", "mode": "submission"}
    )
    assert response.status_code == 200, response.text


# --- The walk-through -------------------------------------------------------------------------


async def test_mvp_steps_end_to_end(
    app_client: httpx.AsyncClient,
    db_session: Session,
    heuristic_llm: HeuristicFakeLLM,
    fake_embeddings,  # noqa: ANN001
    bound_sessions,  # noqa: ANN001
    run_jobs,  # noqa: ANN001
    eval_data: EvalData,
) -> None:
    library = await step_1_library(app_client, eval_data, run_jobs, db_session)
    tender = await step_2_tender(app_client, eval_data, run_jobs)
    draft = await step_3_draft(app_client, db_session, heuristic_llm, tender["questions"])
    review = await step_4_review(
        app_client, db_session, draft["question"], library["newer_reference"]["id"]
    )
    await step_5_export(app_client, tender["tender"]["id"], review["question"])
    await fact_expiry_path(app_client, db_session, tender["tender"]["id"], review["question"]["id"])
    calls = [call.name for call in heuristic_llm.calls]
    assert {
        "classify_document",
        "extract_pairs",
        "chunk_annotate",
        "extract_questions",
        "coverage_judgement",
        "synthesis",
        "entailment",
    } <= set(calls), sorted(set(calls))
