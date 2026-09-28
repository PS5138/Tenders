"""docx export via python-docx.

The document carries the tender name and buyer as its title, one heading per question in the
buyer's order, the question text, and then either the answer prose or the placeholder line.
Review mode writes each sentence as its own run and anchors a Word comment on it listing the
sentence's support status and sources, so a reviewer reads the trace in Word's margin.
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from typing import Any

from docx import Document as new_document
from docx.document import Document
from docx.text.paragraph import Paragraph

from app.db.enums import SegmentKind, SourceType, SupportStatus
from app.db.models import Tender
from app.export import PLACEHOLDER, ExportRow, split_paragraphs

COMMENT_AUTHOR = "Source trace"
COMMENT_INITIALS = "ST"

REVIEW_NOTE = (
    "Internal review copy. Each sentence's sources are attached as comments. Not for submission."
)

SUPPORT_LABELS: dict[str, str] = {
    SupportStatus.SUPPORTED.value: "Supported",
    SupportStatus.WEAK.value: "Weak",
    SupportStatus.UNSUPPORTED.value: "Unsupported",
    SupportStatus.HUMAN_AUTHORED.value: "Human-authored",
    SupportStatus.CONNECTIVE.value: "Connective",
}


def build_docx(tender: Tender, rows: Sequence[ExportRow], *, mode: str) -> bytes:
    """Render ``rows`` (already filtered by the mode's rule) to docx bytes."""
    document = new_document()
    _write_title(document, tender, mode)

    previous_section: str | None = None
    for row in rows:
        if row.section and row.section != previous_section:
            document.add_heading(row.section, level=1)
        previous_section = row.section

        document.add_heading(question_heading(row), level=2)
        question_paragraph = document.add_paragraph()
        question_paragraph.add_run(row.text).italic = True

        if not row.has_answer:
            document.add_paragraph(PLACEHOLDER)
        elif mode == "review" and row.segments:
            _write_segments_with_comments(document, row.segments)
        else:
            _write_prose(document, row.answer_text or "")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def document_title(tender: Tender) -> str:
    """Tender name and buyer, joined with an en dash when both exist."""
    if tender.buyer:
        return f"{tender.name} – {tender.buyer}"
    return tender.name


def question_heading(row: ExportRow) -> str:
    return f"Question {row.number}".strip()


def _write_title(document: Document, tender: Tender, mode: str) -> None:
    document.core_properties.title = document_title(tender)
    if tender.buyer:
        document.core_properties.subject = tender.buyer
    document.add_heading(document_title(tender), level=0)
    if mode == "review":
        document.add_paragraph(REVIEW_NOTE).runs[0].italic = True


def _write_prose(document: Document, text: str) -> None:
    paragraphs = split_paragraphs(text)
    if not paragraphs:
        document.add_paragraph("")
        return
    for paragraph_text in paragraphs:
        document.add_paragraph(paragraph_text)


def _write_segments_with_comments(document: Document, segments: Iterable[dict[str, Any]]) -> None:
    """One docx paragraph per segment ``paragraph`` index, one run per sentence, a comment per
    sentence that has anything to say about its sources."""
    current_paragraph: Paragraph | None = None
    current_index: int | None = None
    for segment in segments:
        paragraph_index = int(segment.get("paragraph") or 0)
        if current_paragraph is None or paragraph_index != current_index:
            current_paragraph = document.add_paragraph()
            current_index = paragraph_index
        elif current_paragraph.runs:
            current_paragraph.add_run(" ")

        run = current_paragraph.add_run(str(segment.get("text", "")))
        comment_text = source_comment(segment)
        if comment_text:
            document.add_comment(
                run, text=comment_text, author=COMMENT_AUTHOR, initials=COMMENT_INITIALS
            )


def source_comment(segment: dict[str, Any]) -> str | None:
    """The comment text for one segment, or None when there is nothing to attach.

    A connective sentence with no sources gets no comment. Every other sentence gets its
    support status, any dispute, and one block per source.
    """
    sources = list(segment.get("sources") or [])
    kind = segment.get("kind", SegmentKind.SUBSTANTIVE.value)
    if not sources and kind == SegmentKind.CONNECTIVE.value:
        return None

    status = str(segment.get("support_status") or "")
    lines = [f"Support: {SUPPORT_LABELS.get(status, status or 'unknown')}"]

    dispute = segment.get("dispute")
    if dispute:
        who = dispute.get("disputed_by") or "a reviewer"
        when = dispute.get("at")
        note = dispute.get("note")
        line = f"Disputed by {who}"
        if when:
            line += f" at {when}"
        if note:
            line += f": {note}"
        lines.append(line)

    if not sources:
        lines.append("No source located for this sentence.")
        return "\n".join(lines)

    for position, source in enumerate(sources, start=1):
        lines.append("")
        lines.extend(_source_lines(position, source))
    return "\n".join(lines)


def _source_lines(position: int, source: dict[str, Any]) -> list[str]:
    source_type = source.get("source_type")
    if source_type == SourceType.HUMAN_ATTESTATION.value:
        who = source.get("attested_by") or "a reviewer"
        line = f"{position}. Attested by {who}"
        if source.get("at"):
            line += f" at {source['at']}"
        if source.get("note"):
            line += f": {source['note']}"
        return [line]

    label = "Fact" if source_type == SourceType.FACT.value else "Source"
    title = source.get("document_title") or "Untitled document"
    descriptors = [
        str(value)
        for value in (source.get("doc_type"), source.get("doc_kind"))
        if value
    ]
    header = f"{position}. {label}: {title}"
    if descriptors:
        header += f" ({', '.join(descriptors)})"
    lines = [header]
    if source.get("effective_date"):
        lines.append(f"Effective {source['effective_date']}")

    location = _location(source.get("locator") or {})
    if location:
        lines.append(f"Location: {location}")

    quote = source.get("quote")
    if quote:
        lines.append(f"“{quote}”")
    return lines


def _location(locator: dict[str, Any]) -> str:
    parts: list[str] = []
    if locator.get("page") is not None:
        parts.append(f"page {locator['page']}")
    if locator.get("table") is not None:
        parts.append(f"table {locator['table']}")
    if locator.get("cell_ref"):
        parts.append(f"row {locator['cell_ref']}")
    heading_path = locator.get("heading_path") or []
    if heading_path:
        parts.append("section " + " > ".join(str(part) for part in heading_path))
    if locator.get("start") is None or locator.get("end") is None:
        parts.append("span not located in this section")
    return ", ".join(parts)


__all__ = [
    "COMMENT_AUTHOR",
    "REVIEW_NOTE",
    "SUPPORT_LABELS",
    "build_docx",
    "document_title",
    "question_heading",
    "source_comment",
]
