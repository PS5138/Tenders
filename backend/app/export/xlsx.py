"""xlsx export via openpyxl: one row per question.

Columns: section, number, question, answer or placeholder, status, word count. The header
sits in row 1 so the sheet filters and sorts in Excel without clean-up; the word count is the
written answer's stored ``word_count`` and is left blank where the placeholder was written.
"""

from __future__ import annotations

import io
from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.db.models import Tender
from app.export import PLACEHOLDER, ExportRow

SHEET_TITLE = "Responses"
HEADERS: tuple[str, ...] = ("Section", "Number", "Question", "Answer", "Status", "Word count")
COLUMN_WIDTHS: tuple[int, ...] = (24, 10, 60, 80, 16, 12)


def build_xlsx(tender: Tender, rows: Sequence[ExportRow], *, mode: str) -> bytes:
    """Render ``rows`` (already filtered by the mode's rule) to xlsx bytes."""
    workbook = Workbook()
    workbook.properties.title = tender.name
    if tender.buyer:
        workbook.properties.subject = tender.buyer
    workbook.properties.description = f"{mode.capitalize()} export"

    sheet = workbook.active
    sheet.title = SHEET_TITLE
    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top")

    for row in rows:
        sheet.append(row_values(row))

    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    for excel_row in sheet.iter_rows(min_row=2):
        for cell in excel_row:
            cell.alignment = wrap if cell.column in (3, 4) else top

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def row_values(row: ExportRow) -> list[object]:
    """The six cell values for one question, in ``HEADERS`` order."""
    return [
        row.section,
        row.number,
        row.text,
        row.answer_text if row.has_answer else PLACEHOLDER,
        row.status_label,
        row.word_count if row.has_answer else None,
    ]


__all__ = ["HEADERS", "SHEET_TITLE", "build_xlsx", "row_values"]
