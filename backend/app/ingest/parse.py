"""Step 1: parse an uploaded file into sections at the plan's granularity and persist them.

Granularity (data model, ``document_sections``):

- Prose is one section per heading-delimited block: the heading line and the paragraphs
  beneath it up to the next heading of any level. Sections never overlap, every character of
  prose belongs to exactly one section, and ``heading_path`` carries the ancestry. The text
  begins with the section's own heading line, except for prose before the first heading, which
  is one section with an empty ``heading_path``.
- Every docx table and every xlsx sheet is one section per row. ``heading_path`` is the
  enclosing prose heading path followed by the header row's cell texts; ``text`` is the row's
  cells in column order joined by the configured cell delimiter. ``table_index`` is set and
  ``cell_ref`` is ``T2:R5`` (docx) or ``Sheet1!5`` (xlsx). Page numbers are null for
  spreadsheets and for docx (python-docx exposes none); the PDF path records them.

docx and xlsx are read with python-docx and openpyxl, which need no models; PDF goes through
Docling with the configured artifacts path and OCR off. Docling is slow to import and is
imported lazily inside the PDF branch only.
"""

from __future__ import annotations

import logging
import re
from collections.abc import Iterable
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db.models import Document, DocumentSection
from app.ingest.types import ParsedSection

logger = logging.getLogger(__name__)

_HEADING_STYLE_RE = re.compile(r"^heading\s*(\d+)$", re.IGNORECASE)
_MARKDOWN_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")


class UnsupportedFileType(ValueError):
    pass


# --- Section builder ----------------------------------------------------------------------------


class _SectionBuilder:
    """Accumulates sections in document order while tracking the heading stack.

    A prose section stays open across tables under the same heading, so paragraphs that follow
    a table still belong to their heading's block; its position in the output is where the
    heading appeared.
    """

    def __init__(self) -> None:
        self.sections: list[ParsedSection] = []
        self._stack: list[tuple[int, str]] = []
        self._open: ParsedSection | None = None
        self._open_lines: list[str] = []
        self._table_count = 0

    @property
    def heading_path(self) -> list[str]:
        return [title for _, title in self._stack]

    def _flush_open(self) -> None:
        if self._open is not None:
            self._open.text = "\n".join(self._open_lines).strip()
        self._open = None
        self._open_lines = []

    def heading(self, level: int, text: str, page: int | None = None) -> None:
        text = _clean_line(text)
        if not text:
            return
        self._flush_open()
        while self._stack and self._stack[-1][0] >= level:
            self._stack.pop()
        self._stack.append((level, text))
        self._open = ParsedSection(
            order_index=len(self.sections),
            heading_path=self.heading_path,
            page_start=page,
            page_end=page,
            table_index=None,
            cell_ref=None,
            text="",
        )
        self._open_lines = [text]
        self.sections.append(self._open)

    def paragraph(self, text: str, page: int | None = None) -> None:
        text = _clean_block(text)
        if not text:
            return
        if self._open is None:
            # Prose before the first heading, or after a table with no heading above it.
            self._open = ParsedSection(
                order_index=len(self.sections),
                heading_path=self.heading_path,
                page_start=page,
                page_end=page,
                table_index=None,
                cell_ref=None,
                text="",
            )
            self._open_lines = []
            self.sections.append(self._open)
        self._open_lines.append(text)
        if page is not None:
            if self._open.page_start is None:
                self._open.page_start = page
            self._open.page_end = page

    def table(
        self,
        rows: Iterable[list[str]],
        *,
        cell_ref: str | None = None,
        page: int | None = None,
        header_from_first_row: bool = True,
    ) -> None:
        """Emit one section per data row. ``cell_ref`` is a format with ``{row}`` for the
        1-based row number within the table; the default is ``T<n>:R<row>``."""
        self._table_count += 1
        table_index = self._table_count
        template = cell_ref or f"T{table_index}:R{{row}}"
        cleaned = [[_clean_cell(cell) for cell in row] for row in rows]
        cleaned = [_strip_trailing_empty(row) for row in cleaned]
        non_empty = [(i, row) for i, row in enumerate(cleaned, start=1) if any(row)]
        if not non_empty:
            return
        header: list[str] = []
        if header_from_first_row and len(non_empty) >= 2:
            _, header = non_empty[0]
            data_rows = non_empty[1:]
        else:
            data_rows = non_empty
        base_path = self.heading_path
        delimiter = get_settings().table_cell_delimiter
        for row_number, row in data_rows:
            self.sections.append(
                ParsedSection(
                    order_index=len(self.sections),
                    heading_path=[*base_path, *header],
                    page_start=page,
                    page_end=page,
                    table_index=table_index,
                    cell_ref=template.format(row=row_number),
                    text=delimiter.join(row),
                )
            )

    def finish(self) -> list[ParsedSection]:
        self._flush_open()
        kept = [section for section in self.sections if section.text.strip()]
        for index, section in enumerate(kept):
            section.order_index = index
        return kept


def _clean_line(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def _clean_block(text: str) -> str:
    lines = [line.rstrip() for line in (text or "").replace("\r", "").split("\n")]
    return "\n".join(lines).strip()


def _clean_cell(text: str | None) -> str:
    if text is None:
        return ""
    return _clean_block(str(text)).replace("\n", " ")


def _strip_trailing_empty(row: list[str]) -> list[str]:
    end = len(row)
    while end > 0 and not row[end - 1]:
        end -= 1
    return row[:end]


# --- docx ---------------------------------------------------------------------------------------


def _docx_heading_level(paragraph) -> int | None:  # noqa: ANN001
    style = paragraph.style
    name = (style.name if style is not None else "") or ""
    if name.lower() == "title":
        return 1
    match = _HEADING_STYLE_RE.match(name)
    if match:
        return int(match.group(1))
    # Fall back to the outline level a custom style may declare.
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        outline = p_pr.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl"
        )
        if outline is not None:
            value = outline.get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
            )
            if value is not None and value.isdigit() and int(value) < 9:
                return int(value) + 1
    return None


def _docx_row_cells(row) -> list[str]:  # noqa: ANN001
    """Cell texts in column order; a horizontally merged cell is reported once."""
    cells: list[str] = []
    seen_tc = None
    for cell in row.cells:
        if cell._tc is seen_tc:
            continue
        seen_tc = cell._tc
        cells.append(cell.text)
    return cells


def _parse_docx(path: Path) -> list[ParsedSection]:
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    docx_document = DocxDocument(str(path))
    builder = _SectionBuilder()
    body = docx_document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, docx_document)
            level = _docx_heading_level(paragraph)
            if level is not None:
                builder.heading(level, paragraph.text)
            else:
                builder.paragraph(paragraph.text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, docx_document)
            builder.table(_docx_row_cells(row) for row in table.rows)
    return builder.finish()


# --- xlsx ---------------------------------------------------------------------------------------


def _cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _parse_xlsx(path: Path) -> list[ParsedSection]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    builder = _SectionBuilder()
    delimiter = get_settings().table_cell_delimiter
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
            header: list[str] | None = None
            base_path = [sheet.title]
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = _strip_trailing_empty([_clean_cell(_cell_value(v)) for v in row])
                if not any(cells):
                    continue
                if header is None:
                    header = cells
                    continue
                builder.sections.append(
                    ParsedSection(
                        order_index=len(builder.sections),
                        heading_path=[*base_path, *header],
                        page_start=None,
                        page_end=None,
                        table_index=sheet_index,
                        cell_ref=f"{sheet.title}!{row_number}",
                        text=delimiter.join(cells),
                    )
                )
            if header is not None and not any(
                s.table_index == sheet_index for s in builder.sections
            ):
                # A sheet with a single non-empty row: that row is data, not a header.
                builder.sections.append(
                    ParsedSection(
                        order_index=len(builder.sections),
                        heading_path=base_path,
                        page_start=None,
                        page_end=None,
                        table_index=sheet_index,
                        cell_ref=f"{sheet.title}!1",
                        text=delimiter.join(header),
                    )
                )
    finally:
        workbook.close()
    return builder.finish()


# --- PDF (Docling) ------------------------------------------------------------------------------


def _docling_page(item) -> int | None:  # noqa: ANN001
    prov = getattr(item, "prov", None) or []
    if not prov:
        return None
    page = getattr(prov[0], "page_no", None)
    return int(page) if page is not None else None


def _docling_table_rows(item) -> list[list[str]]:  # noqa: ANN001
    data = getattr(item, "data", None)
    if data is None:
        return []
    grid = getattr(data, "grid", None)
    if grid:
        return [[(getattr(cell, "text", "") or "") for cell in row] for row in grid]
    rows: dict[int, dict[int, str]] = {}
    for cell in getattr(data, "table_cells", []) or []:
        rows.setdefault(cell.start_row_offset_idx, {})[cell.start_col_offset_idx] = cell.text or ""
    return [
        [row.get(col, "") for col in range(max(row) + 1)] if row else []
        for _, row in sorted(rows.items())
    ]


def _parse_pdf(path: Path) -> list[ParsedSection]:
    # Docling is heavy; import it here so the API process never pays for it.
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docling_core.types.doc import (
        ListItem,
        SectionHeaderItem,
        TableItem,
        TextItem,
        TitleItem,
    )

    settings = get_settings()
    options = PdfPipelineOptions(do_ocr=False)
    if settings.docling_artifacts_path is not None:
        options.artifacts_path = str(settings.docling_artifacts_path)
    converter = DocumentConverter(
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)}
    )
    result = converter.convert(str(path))
    document = result.document

    builder = _SectionBuilder()
    for item, _level in document.iterate_items():
        page = _docling_page(item)
        if isinstance(item, TitleItem):
            builder.heading(1, item.text, page)
        elif isinstance(item, SectionHeaderItem):
            builder.heading(int(getattr(item, "level", 1) or 1), item.text, page)
        elif isinstance(item, TableItem):
            builder.table(_docling_table_rows(item), page=page)
        elif isinstance(item, ListItem | TextItem):
            builder.paragraph(item.text, page)
    return builder.finish()


# --- Plain text and Markdown --------------------------------------------------------------------


def _parse_text(path: Path) -> list[ParsedSection]:
    builder = _SectionBuilder()
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            builder.paragraph("\n".join(buffer))
            buffer.clear()

    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        match = _MARKDOWN_HEADING_RE.match(line)
        if match:
            flush()
            builder.heading(len(match.group(1)), match.group(2))
        elif line.strip():
            buffer.append(line.rstrip())
        else:
            flush()
    flush()
    return builder.finish()


# --- Public API ---------------------------------------------------------------------------------

_PARSERS = {
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".xlsm": _parse_xlsx,
    ".pdf": _parse_pdf,
    ".txt": _parse_text,
    ".md": _parse_text,
}


def parse_document(path: Path, filename: str) -> list[ParsedSection]:
    """Parse the file at ``path`` (``filename`` decides the format) into ordered sections."""
    suffix = Path(filename).suffix.lower() or Path(path).suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise UnsupportedFileType(
            f"Unsupported file type '{suffix or 'none'}'. Upload a .docx, .xlsx or .pdf file."
        )
    sections = parser(Path(path))
    logger.info("parsed %s into %d section(s)", filename, len(sections))
    return sections


def persist_sections(
    session: Session, document: Document, parsed: list[ParsedSection]
) -> list[DocumentSection]:
    """Write the sections once. Sections are never edited or deleted, so a re-run of the parse
    stage returns the rows already stored instead of writing duplicates."""
    existing = session.scalars(
        select(DocumentSection)
        .where(DocumentSection.document_id == document.id)
        .order_by(DocumentSection.order_index)
    ).all()
    if existing:
        return list(existing)
    rows = [
        DocumentSection(
            org_id=document.org_id,
            document_id=document.id,
            order_index=section.order_index,
            heading_path=list(section.heading_path),
            page_start=section.page_start,
            page_end=section.page_end,
            table_index=section.table_index,
            cell_ref=section.cell_ref,
            text=section.text,
        )
        for section in parsed
    ]
    session.add_all(rows)
    session.flush()
    return rows


__all__ = ["UnsupportedFileType", "parse_document", "persist_sections"]
